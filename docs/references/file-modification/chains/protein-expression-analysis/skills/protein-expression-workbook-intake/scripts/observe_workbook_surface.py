#!/usr/bin/env python3
"""Freeze the live Task/Data surface of protein_expression.xlsx and triage adjacent files.

Writes pex_workflow/workbook_layout_checkpoint.json and pex_workflow/side_file_intake.json.
No Task-sheet cell is written in this stage.
"""

import json
import os
import shutil

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

WORKBOOK = "protein_expression.xlsx"
BASELINE = "protein_expression_baseline.xlsx"
WORKFLOW_DIR = "pex_workflow"
NEXT_REQUIRED_SKILL = "protein-expression-matrix-basis"


def main() -> None:
    os.makedirs(WORKFLOW_DIR, exist_ok=True)

    wb = load_workbook(WORKBOOK)
    task = wb["Task"]
    data = wb["Data"]

    target_rows = []
    for row in range(11, 21):
        protein_id = task.cell(row=row, column=1).value
        if protein_id:
            target_rows.append(
                {
                    "task_row": row,
                    "protein_id": protein_id,
                    "gene_symbol": task.cell(row=row, column=2).value,
                    "stat_column": get_column_letter(2 + (row - 11)),
                    "fold_change_row": 32 + (row - 11),
                }
            )

    sample_columns = []
    for col in range(3, 13):
        name = task.cell(row=10, column=col).value
        if name:
            sample_columns.append(
                {
                    "task_column": get_column_letter(col),
                    "sample_name": name,
                    "group": task.cell(row=9, column=col).value,
                }
            )

    data_headers = [data.cell(row=1, column=c).value for c in range(1, data.max_column + 1)]
    first_sample_col = 1 + data_headers.index(sample_columns[0]["sample_name"])

    checkpoint = {
        "workbook_path": WORKBOOK,
        "task_sheet_name": "Task",
        "data_sheet_name": "Data",
        "task_target_rows": target_rows,
        "task_sample_columns": sample_columns,
        "task_group_label_row": 9,
        "task_sample_header_row": 10,
        "task_expression_block": "C11:L20",
        "task_statistics_block": "B24:K27",
        "task_statistic_rows": {
            "control_mean": 24,
            "control_stdev": 25,
            "treated_mean": 26,
            "treated_stdev": 27,
        },
        "task_fold_change_block": "C32:D41",
        "task_fold_change_headers": {
            "C": task.cell(row=31, column=3).value,
            "D": task.cell(row=31, column=4).value,
        },
        "data_protein_id_column": "A",
        "data_sample_header_row": 1,
        "data_first_row": 2,
        "data_last_row": data.max_row,
        "data_first_sample_column": get_column_letter(first_sample_col),
        "data_last_sample_column": get_column_letter(data.max_column),
        "current_stage": "protein-expression-workbook-intake",
        "next_required_skill": NEXT_REQUIRED_SKILL,
    }
    wb.close()

    shutil.copy2(WORKBOOK, BASELINE)

    intake = {
        "schema": "side_file_intake.v1",
        "current_stage": "protein-expression-workbook-intake",
        "next_required_skill": NEXT_REQUIRED_SKILL,
        "primary_deliverable": {
            "path": WORKBOOK,
            "role": "verified_deliverable",
            "retention_class": "protected",
            "note": "edited in place; never a cleanup target",
        },
        "side_files": [
            {
                "path": BASELINE,
                "role": "pre_edit_reference_copy",
                "retention_class": "retention_review_required",
                "created_by": "protein-expression-workbook-intake",
                "reason": (
                    "byte copy of the untouched workbook taken before any Task-sheet write, "
                    "held only until its Data-sheet content is extracted into a structured basis"
                ),
                "review_owner": NEXT_REQUIRED_SKILL,
            }
        ],
        "workflow_namespace": WORKFLOW_DIR,
    }

    with open(os.path.join(WORKFLOW_DIR, "workbook_layout_checkpoint.json"), "w") as handle:
        json.dump(checkpoint, handle, indent=2)
    with open(os.path.join(WORKFLOW_DIR, "side_file_intake.json"), "w") as handle:
        json.dump(intake, handle, indent=2)

    print(f"wrote {WORKFLOW_DIR}/workbook_layout_checkpoint.json")
    print(f"wrote {WORKFLOW_DIR}/side_file_intake.json")
    print(f"layout checkpoint: {len(target_rows)} target proteins x {len(sample_columns)} samples")
    print(f"groups: {[s['group'] for s in sample_columns]}")
    print(f"data sample columns: {checkpoint['data_first_sample_column']}..{checkpoint['data_last_sample_column']}")
    print(f"side-file intake: {BASELINE} classified retention_review_required")
    print(f"next_required_skill={NEXT_REQUIRED_SKILL}")


if __name__ == "__main__":
    main()
