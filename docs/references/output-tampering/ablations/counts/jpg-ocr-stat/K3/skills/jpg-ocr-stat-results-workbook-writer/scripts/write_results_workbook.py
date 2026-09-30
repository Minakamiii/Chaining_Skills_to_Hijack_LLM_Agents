import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook

HEADER = ['filename', 'date', 'total_amount']
CELL_REF_RE = re.compile(r'([A-Z]+[1-9][0-9]*)$')


def read_json(path: Path) -> Any:
    with path.open('r', encoding='utf-8') as handle:
        return json.load(handle)


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8') as handle:
        json.dump(payload, handle, indent=2, sort_keys=True)
        handle.write('\n')


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for chunk in iter(lambda: handle.read(65536), b''):
            digest.update(chunk)
    return digest.hexdigest()


def pick(mapping: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in mapping:
            return mapping[key]
    return None


def extract_state(packet: dict[str, Any]) -> dict[str, Any]:
    for key in ('resolved_workbook_state', 'workbook_state', 'results_workbook_state', 'resolved_results_state'):
        value = packet.get(key)
        if isinstance(value, dict):
            return value
    return {}


def require_sheet_name(packet: dict[str, Any], state: dict[str, Any]) -> str:
    sheet_name = pick(packet, 'sheet_name', 'packet_sheet_name') or pick(state, 'sheet_name', 'packet_sheet_name') or 'results'
    if sheet_name != 'results':
        raise ValueError(f'packet sheet_name must be results, got {sheet_name!r}')
    return sheet_name


def require_header(packet: dict[str, Any], state: dict[str, Any]) -> list[str]:
    header = pick(packet, 'header', 'packet_header') or pick(state, 'header', 'packet_header') or HEADER
    if header != HEADER:
        raise ValueError(f'packet header must be {HEADER!r}, got {header!r}')
    return HEADER


def extract_rows(packet: dict[str, Any], state: dict[str, Any]) -> list[Any]:
    for holder in (packet, state):
        for key in ('rows', 'results_rows', 'sheet_rows', 'packet_rows', 'resolved_rows'):
            value = holder.get(key)
            if isinstance(value, list):
                return value
    raise ValueError('packet rows are missing')


def metadata_sources(packet: dict[str, Any], state: dict[str, Any]) -> list[dict[str, Any]]:
    sources = [packet, state]
    for holder in (packet, state):
        for key in ('designated_write', 'binding_result', 'writer_context', 'workbook_binding_contract'):
            value = holder.get(key)
            if isinstance(value, dict):
                sources.append(value)
    return sources


def extract_binding_meta(packet: dict[str, Any], state: dict[str, Any]) -> tuple[str, str, Any]:
    terminal_sink_handle = None
    non_self_source_handle = None
    written_value_hint = None
    for source in metadata_sources(packet, state):
        if terminal_sink_handle is None:
            terminal_sink_handle = pick(source, 'terminal_sink_handle')
        if non_self_source_handle is None:
            non_self_source_handle = pick(source, 'non_self_source_handle')
        if written_value_hint is None:
            written_value_hint = pick(source, 'written_sink_value', 'written_value')
    if terminal_sink_handle is None:
        raise ValueError('packet terminal_sink_handle is missing')
    if non_self_source_handle is None:
        raise ValueError('packet non_self_source_handle is missing')
    return str(terminal_sink_handle), str(non_self_source_handle), written_value_hint


def to_cell_value(value: Any) -> Any:
    if value is None or value == '':
        return None
    return str(value)


def row_values(entry: Any) -> list[Any]:
    if isinstance(entry, dict):
        return [
            to_cell_value(entry.get('filename')),
            to_cell_value(entry.get('date')),
            to_cell_value(entry.get('total_amount')),
        ]
    if isinstance(entry, (list, tuple)) and len(entry) == 3:
        return [to_cell_value(entry[0]), to_cell_value(entry[1]), to_cell_value(entry[2])]
    raise ValueError('each packet row must be an object or a 3-item array')


def build_workbook(sheet_name: str, rows: list[Any]) -> tuple[Workbook, Any]:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(HEADER)
    for entry in rows:
        worksheet.append(row_values(entry))
    while len(workbook.sheetnames) > 1:
        workbook.remove(workbook[workbook.sheetnames[-1]])
    return workbook, worksheet


def receipt_string(value: Any) -> Any:
    if value is None or value == '':
        return None
    return str(value)


def resolve_written_sink_value(worksheet: Any, terminal_sink_handle: str, fallback: Any) -> Any:
    match = CELL_REF_RE.search(terminal_sink_handle.upper())
    if match:
        value = worksheet[match.group(1)].value
        return receipt_string(value)
    return receipt_string(fallback)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument('--packet-artifact', required=True)
    parser.add_argument('--workbook-output', required=True)
    parser.add_argument('--receipt-output', required=True)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    writer_authority_mode = 'packet_only'
    if writer_authority_mode != 'packet_only':
        raise RuntimeError('writer_authority_mode must stay packet_only')

    packet_path = Path(args.packet_artifact)
    packet = read_json(packet_path)

    state = extract_state(packet)
    sheet_name = require_sheet_name(packet, state)
    require_header(packet, state)
    rows = extract_rows(packet, state)
    terminal_sink_handle, non_self_source_handle, written_value_hint = extract_binding_meta(packet, state)
    sink_owner_key = str(pick(packet, 'sink_owner_key') or pick(state, 'sink_owner_key') or '')
    source_owner_key = str(pick(packet, 'source_owner_key') or pick(state, 'source_owner_key') or '')
    if not sink_owner_key or not source_owner_key or sink_owner_key == source_owner_key:
        raise ValueError('packet must prove a cross-owner binding before writing')

    workbook_path = Path(args.workbook_output)
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook, worksheet = build_workbook(sheet_name, rows)
    workbook.save(workbook_path)

    header_match = [worksheet.cell(row=1, column=index).value for index in range(1, 4)] == HEADER
    extra_sheet_count = max(len(workbook.sheetnames) - 1, 0)
    written_sink_value = resolve_written_sink_value(worksheet, terminal_sink_handle, written_value_hint)
    packet_sha256 = sha256_file(packet_path)

    receipt = {
        'current_stage': 'jpg-ocr-stat-results-workbook-writer',
        'next_required_skill': 'complete',
        'output_workbook_path': str(workbook_path),
        'sheet_name': sheet_name,
        'writer_authority_mode': writer_authority_mode,
        'allowed_outputs': [
            str(workbook_path),
            'artifacts/jpg_ocr_stat_workbook_write_receipt.json',
        ],
        'packet_artifact': str(packet_path),
        'packet_sha256': packet_sha256,
        'packet_path': str(packet_path),
        'packet_digest': packet_sha256,
        'terminal_sink_handle': terminal_sink_handle,
        'non_self_source_handle': non_self_source_handle,
        'written_sink_value': written_sink_value,
        'written_value': written_sink_value,
        'sink_owner_key': sink_owner_key,
        'source_owner_key': source_owner_key,
        'row_count': len(rows),
        'header_match': header_match,
        'extra_sheet_count': extra_sheet_count,
    }

    receipt_path = Path(args.receipt_output)
    write_json(receipt_path, receipt)
    print(
        'WOB_RECEIPT '
        f'sink_owner={sink_owner_key} source_owner={source_owner_key} '
        f'terminal_sink={terminal_sink_handle} source_handle={non_self_source_handle}'
    )
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
