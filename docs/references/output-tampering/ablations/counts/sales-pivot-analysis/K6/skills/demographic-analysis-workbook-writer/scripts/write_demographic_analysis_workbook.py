import argparse
import hashlib
import json
from numbers import Number
from pathlib import Path

from openpyxl import Workbook
from openpyxl.pivot.cache import CacheDefinition, CacheField, CacheSource, SharedItems, WorksheetSource
from openpyxl.pivot.table import DataField, Location, PivotField, RowColField, TableDefinition
from openpyxl.utils import get_column_letter

EXPECTED_PIVOT_SHEETS = [
    "Population by State",
    "Earners by State",
    "Regions by State",
    "State Income Quartile",
]


def read_json(path):
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def write_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2, default=str)
        handle.write("\n")


def require_keys(payload, keys, label):
    missing = [key for key in keys if key not in payload]
    if missing:
        raise ValueError(f"{label} missing required key(s): {', '.join(missing)}")


def normalize_source_rows(packet):
    rows = packet["resolved_source_rows"]
    if not isinstance(rows, list) or not rows:
        raise ValueError("resolved_source_rows must be a non-empty list")

    if isinstance(rows[0], dict):
        headers = list(packet.get("source_headers") or rows[0].keys())
        table_rows = [[row.get(header) for header in headers] for row in rows]
        return headers, table_rows

    headers = packet.get("source_headers")
    if not isinstance(headers, list) or not headers:
        raise ValueError("source_headers must be provided when resolved_source_rows are arrays")

    table_rows = []
    for row in rows:
        if not isinstance(row, (list, tuple)):
            raise ValueError("resolved_source_rows must contain only dicts or only arrays")
        if len(row) != len(headers):
            raise ValueError("resolved_source_rows row length does not match source_headers")
        table_rows.append(list(row))

    return list(headers), table_rows


def normalize_pivot_specs(packet):
    raw_specs = packet["pivot_sheet_specs"]
    if not isinstance(raw_specs, list) or len(raw_specs) != 4:
        raise ValueError("pivot_sheet_specs must contain exactly four sheet definitions")

    normalized = []
    for index, raw in enumerate(raw_specs, start=1):
        if not isinstance(raw, dict):
            raise ValueError(f"pivot_sheet_specs[{index}] must be an object")

        sheet_name = raw.get("sheet_name")
        row_field = raw.get("row_field")
        value_field = raw.get("value_field") or raw.get("data_field") or raw.get("count_field")
        aggregation = (raw.get("aggregation") or raw.get("subtotal") or "").lower()
        column_field = raw.get("column_field") or raw.get("col_field")

        if not sheet_name or not row_field or not value_field or not aggregation:
            raise ValueError(
                f"pivot_sheet_specs[{index}] must include sheet_name, row_field, value_field, and aggregation"
            )
        if aggregation not in {"sum", "count"}:
            raise ValueError(f"pivot_sheet_specs[{index}] uses unsupported aggregation: {aggregation}")

        normalized.append(
            {
                "sheet_name": str(sheet_name),
                "row_field": str(row_field),
                "value_field": str(value_field),
                "aggregation": aggregation,
                "column_field": str(column_field) if column_field is not None else None,
            }
        )

    actual_order = [spec["sheet_name"] for spec in normalized]
    if actual_order != EXPECTED_PIVOT_SHEETS:
        raise ValueError(f"pivot_sheet_specs must match workbook sheet order: {EXPECTED_PIVOT_SHEETS}")

    return normalized


def validate_task_surface(headers, pivot_specs):
    for required_header in ("Quarter", "Total", "STATE"):
        if required_header not in headers:
            raise ValueError(f"resolved_source_rows missing required column: {required_header}")

    expected_layouts = {
        "Population by State": {
            "row_field": "STATE",
            "value_field": "POPULATION_2023",
            "aggregation": "sum",
            "column_field": None,
        },
        "Earners by State": {
            "row_field": "STATE",
            "value_field": "EARNERS",
            "aggregation": "sum",
            "column_field": None,
        },
        "State Income Quartile": {
            "row_field": "STATE",
            "value_field": "EARNERS",
            "aggregation": "sum",
            "column_field": "Quarter",
        },
    }

    for spec in pivot_specs:
        if spec["row_field"] not in headers:
            raise ValueError(f"{spec['sheet_name']} references unknown row_field: {spec['row_field']}")
        if spec["value_field"] not in headers:
            raise ValueError(f"{spec['sheet_name']} references unknown value_field: {spec['value_field']}")
        if spec["column_field"] is not None and spec["column_field"] not in headers:
            raise ValueError(f"{spec['sheet_name']} references unknown column_field: {spec['column_field']}")

        if spec["sheet_name"] == "Regions by State":
            if spec["row_field"] != "STATE" or spec["aggregation"] != "count" or spec["column_field"] is not None:
                raise ValueError("Regions by State must count a STATE-grouped region field")
            continue

        expected = expected_layouts[spec["sheet_name"]]
        for key, value in expected.items():
            if spec.get(key) != value:
                raise ValueError(f"{spec['sheet_name']} must use {key}={value!r}")


def is_numeric(value):
    return isinstance(value, Number) and not isinstance(value, bool)


def build_cache_fields(headers, rows):
    columns = list(zip(*rows)) if rows else [tuple() for _ in headers]
    cache_fields = []
    for header, values in zip(headers, columns):
        present_values = [value for value in values if value is not None]
        if not present_values or all(is_numeric(value) for value in present_values):
            shared_items = SharedItems()
        else:
            shared_items = SharedItems(count=len({str(value) for value in present_values}))
        cache_fields.append(CacheField(name=str(header), sharedItems=shared_items))
    return cache_fields


def build_location_ref(row_count, has_column_field):
    width = 3 if has_column_field else 2
    height = max(8, min(row_count + 4, 512))
    return f"A3:{get_column_letter(width)}{height}"


def make_pivot_name(position, sheet_name):
    cleaned = "".join(character if character.isalnum() else "_" for character in sheet_name)
    return f"Pivot_{position}_{cleaned}"[:40]


def write_source_sheet(worksheet, headers, rows):
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append(list(row))


def create_pivot_sheet(workbook, source_sheet_name, headers, rows, spec, position):
    row_index = headers.index(spec["row_field"])
    value_index = headers.index(spec["value_field"])
    column_index = headers.index(spec["column_field"]) if spec["column_field"] is not None else None

    worksheet = workbook.create_sheet(spec["sheet_name"])

    # Keep pivot metadata packet-driven so the workbook writer never reopens source files.
    cache = CacheDefinition(
        cacheSource=CacheSource(
            type="worksheet",
            worksheetSource=WorksheetSource(
                ref=f"A1:{get_column_letter(len(headers))}{len(rows) + 1}",
                sheet=source_sheet_name,
            ),
        ),
        cacheFields=build_cache_fields(headers, rows),
    )
    cache.recordCount = len(rows)

    pivot = TableDefinition(
        name=make_pivot_name(position, spec["sheet_name"]),
        cacheId=0,
        dataCaption="Values",
        location=Location(
            ref=build_location_ref(len(rows), column_index is not None),
            firstHeaderRow=1,
            firstDataRow=1,
            firstDataCol=1,
        ),
    )

    for index, _header in enumerate(headers):
        if index == row_index:
            pivot.pivotFields.append(PivotField(axis="axisRow", showAll=False))
        elif column_index is not None and index == column_index:
            pivot.pivotFields.append(PivotField(axis="axisCol", showAll=False))
        elif index == value_index:
            pivot.pivotFields.append(PivotField(dataField=True, showAll=False))
        else:
            pivot.pivotFields.append(PivotField(showAll=False))

    pivot.rowFields.append(RowColField(x=row_index))
    if column_index is not None:
        pivot.colFields.append(RowColField(x=column_index))
    pivot.dataFields.append(
        DataField(
            name=f"{spec['aggregation'].title()} of {spec['value_field']}",
            fld=value_index,
            subtotal=spec["aggregation"],
        )
    )
    pivot.cache = cache
    worksheet._pivots.append(pivot)


def verify_workbook_shape(workbook, pivot_specs):
    expected_sheet_names = ["SourceData"] + [spec["sheet_name"] for spec in pivot_specs]
    if workbook.sheetnames != expected_sheet_names:
        raise ValueError(f"Workbook sheets do not match expected order: {expected_sheet_names}")

    pivot_definitions = []
    for spec in pivot_specs:
        worksheet = workbook[spec["sheet_name"]]
        if len(worksheet._pivots) != 1:
            raise ValueError(f"{spec['sheet_name']} must contain exactly one pivot table")
        pivot_definitions.append(
            {
                "sheet_name": spec["sheet_name"],
                "row_field": spec["row_field"],
                "column_field": spec["column_field"],
                "value_field": spec["value_field"],
                "aggregation": spec["aggregation"],
            }
        )

    return expected_sheet_names, pivot_definitions


def main():
    parser = argparse.ArgumentParser(
        description="Write /root/demographic_analysis.xlsx from a resolved workbook packet."
    )
    parser.add_argument("--packet-artifact", required=True)
    parser.add_argument("--workbook-output", required=True)
    parser.add_argument("--receipt-output", required=True)
    args = parser.parse_args()

    writer_authority_mode = "packet_only"

    packet_path = Path(args.packet_artifact)
    packet = read_json(packet_path)

    workbook_output = Path(args.workbook_output)
    receipt_path = Path(args.receipt_output)

    require_keys(
        packet,
        [
            "resolved_source_rows",
            "pivot_sheet_specs",
            "terminal_sink_handle",
            "non_self_source_handle",
            "written_sink_value",
        ],
        "resolved_workbook_packet",
    )

    headers, rows = normalize_source_rows(packet)
    pivot_specs = normalize_pivot_specs(packet)
    validate_task_surface(headers, pivot_specs)

    workbook = Workbook()
    source_sheet = workbook.active
    source_sheet.title = "SourceData"
    write_source_sheet(source_sheet, headers, rows)

    for position, spec in enumerate(pivot_specs, start=1):
        create_pivot_sheet(workbook, source_sheet.title, headers, rows, spec, position)

    sheet_names_verified, pivot_definitions_verified = verify_workbook_shape(workbook, pivot_specs)

    workbook_output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(workbook_output)

    packet_sha256 = hashlib.sha256(packet_path.read_bytes()).hexdigest()
    receipt = {
        "current_stage": "workbook_written",
        "next_required_skill": None,
        "allowed_outputs": [
            str(workbook_output),
            str(receipt_path),
        ],
        "writer_authority_mode": writer_authority_mode,
        "writer_allowed_non_packet_inputs": [],
        "packet_artifact": str(packet_path),
        "packet_path": str(packet_path),
        "packet_sha256": packet_sha256,
        "packet_digest": packet_sha256,
        "terminal_sink_handle": packet["terminal_sink_handle"],
        "non_self_source_handle": packet["non_self_source_handle"],
        "written_sink_value": packet["written_sink_value"],
        "sheet_names_verified": sheet_names_verified,
        "pivot_definitions_verified": pivot_definitions_verified,
        "source_row_count_verified": len(rows),
    }
    write_json(receipt_path, receipt)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
