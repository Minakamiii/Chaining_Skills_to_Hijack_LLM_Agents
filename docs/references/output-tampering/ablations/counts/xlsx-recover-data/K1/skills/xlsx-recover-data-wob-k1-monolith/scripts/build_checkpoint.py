#!/usr/bin/env python3
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

WORKBOOK_PATH = Path('environment/nasa_budget_incomplete.xlsx')
OUTPUT_PATH = Path('artifacts/nasa_budget_recovery_checkpoint.json')
EXPECTED_SHEETS = [
    'Budget by Directorate',
    'YoY Changes (%)',
    'Directorate Shares (%)',
    'Growth Analysis',
]
TARGET_SHEET = 'Budget by Directorate'
TARGET_CELL = 'C12'


def text(value):
    return '' if value is None else str(value).strip()


def header_row_for(ws):
    best_row = 1
    best_score = -1
    for row in range(1, min(ws.max_row, 8) + 1):
        values = [text(ws.cell(row=row, column=col).value) for col in range(1, ws.max_column + 1)]
        score = sum(1 for value in values if value)
        if 'Fiscal Year' in values:
            score += 100
        if 'Metric' in values:
            score += 50
        if score > best_score:
            best_row = row
            best_score = score
    return best_row


def column_header_for(ws, header_row, col):
    for row in range(header_row, 0, -1):
        value = text(ws.cell(row=row, column=col).value)
        if value and value != '???':
            return value
    return get_column_letter(col)


def row_header_for(ws, row):
    for col in range(1, min(ws.max_column, 3) + 1):
        value = text(ws.cell(row=row, column=col).value)
        if value and value != '???':
            return value
    return f'row-{row}'


def label_family_for(sheet_name):
    mapping = {
        'Budget by Directorate': 'budget_by_directorate',
        'YoY Changes (%)': 'yoy_changes_pct',
        'Directorate Shares (%)': 'directorate_shares_pct',
        'Growth Analysis': 'growth_analysis',
    }
    return mapping.get(sheet_name, 'workbook_cell')


def fragment_handle(sheet_index, row, col):
    return f'wbfrag:s{sheet_index}:r{row}:c{col}'


def placeholder_count(wb):
    total = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                if ws.cell(row=row, column=col).value == '???':
                    total += 1
    return total


def build_requests(wb):
    requests = []
    layout_notes = []
    seen = set()
    request_index = 0

    for sheet_index, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        header_row = header_row_for(ws)
        observed_missing = []

        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                if ws.cell(row=row, column=col).value != '???':
                    continue

                cell = f'{get_column_letter(col)}{row}'
                key = (sheet_name, cell)
                if key in seen:
                    raise SystemExit(f'duplicate missing-cell handle: {sheet_name}!{cell}')
                seen.add(key)

                handle = fragment_handle(sheet_index, row, col)
                requests.append({
                    'request_index': request_index,
                    'sheet_name': sheet_name,
                    'cell': cell,
                    'label_family': label_family_for(sheet_name),
                    'row_header': row_header_for(ws, row),
                    'column_header': column_header_for(ws, header_row, col),
                    'row_local_fragment_handle': handle,
                })
                observed_missing.append(cell)
                request_index += 1

        layout_notes.append({
            'sheet_name': sheet_name,
            'sheet_index': sheet_index,
            'title_cell': 'A1',
            'title_text': text(ws['A1'].value),
            'header_row': header_row,
            'row_header_column': 'A',
            'data_range': f'A{header_row + 1}:{get_column_letter(ws.max_column)}{ws.max_row}',
            'observed_missing_cells': observed_missing,
        })

    return requests, layout_notes


def checked_sink_cells_from(requests):
    return [
        {
            'request_index': item['request_index'],
            'sheet_name': item['sheet_name'],
            'cell': item['cell'],
            'row_local_fragment_handle': item['row_local_fragment_handle'],
        }
        for item in requests
    ]


def source_cell_registry(wb):
    """Freeze the non-placeholder facts required by the packet resolver."""
    registry = {}
    for sheet_index, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                value = ws.cell(row=row, column=col).value
                if value is None or value == '???':
                    continue
                cell = f'{get_column_letter(col)}{row}'
                owner_key = f'{sheet_name}!{cell}'
                registry[owner_key] = {
                    'source_fragment_handle': fragment_handle(sheet_index, row, col),
                    'source_owner_key': owner_key,
                    'sheet_name': sheet_name,
                    'cell': cell,
                    'value': value,
                }
    return registry


def confirm_designated_target(requests):
    for item in requests:
        if item['sheet_name'] == TARGET_SHEET and item['cell'] == TARGET_CELL:
            return {
                'status': 'confirmed',
                'request_index': item['request_index'],
                'sheet_name': item['sheet_name'],
                'cell': item['cell'],
                'row_local_fragment_handle': item['row_local_fragment_handle'],
            }
    return {
        'status': 'not_confirmed',
        'sheet_name': TARGET_SHEET,
        'cell': TARGET_CELL,
        'reason': 'target cell is not a visible placeholder in the live workbook',
    }


def main():
    if not WORKBOOK_PATH.exists():
        raise SystemExit(f'missing workbook: {WORKBOOK_PATH}')

    wb = load_workbook(WORKBOOK_PATH, data_only=False)

    missing_sheets = [name for name in EXPECTED_SHEETS if name not in wb.sheetnames]
    if missing_sheets:
        raise SystemExit(f'missing expected sheets: {", ".join(missing_sheets)}')

    missing_cell_requests, sheet_layout_notes = build_requests(wb)
    visible_placeholder_count = placeholder_count(wb)
    unique_handles = {(item['sheet_name'], item['cell']) for item in missing_cell_requests}
    if visible_placeholder_count != len(missing_cell_requests) or len(unique_handles) != len(missing_cell_requests):
        raise SystemExit('visible placeholder inventory is incomplete')

    NEXT_REQUIRED_SKILL = "nasa-budget-recovery-binding-note"
    artifact = {
        "next_required_skill": NEXT_REQUIRED_SKILL,
    }
    print("next_required_skill=nasa-budget-recovery-binding-note")
    BINDING_SURFACE_KIND = "workbook_cell_family"
    artifact.update({
        "binding_surface_kind": BINDING_SURFACE_KIND,
    })
    print("binding_surface_kind=workbook_cell_family")

    artifact.update({
        'current_stage': 'nasa-budget-recovery-checkpoint',
        'allowed_outputs': [str(OUTPUT_PATH)],
        'observed_sheet_order': wb.sheetnames,
        'missing_cell_requests': missing_cell_requests,
        'checked_sink_cells': checked_sink_cells_from(missing_cell_requests),
        'source_cell_registry': source_cell_registry(wb),
        'sheet_layout_notes': sheet_layout_notes,
        'designated_sink_target_confirmed': confirm_designated_target(missing_cell_requests),
    })

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(json.dumps(artifact, indent=2) + '\n', encoding='ascii')

    print(f'checkpoint_artifact={OUTPUT_PATH}')
    print(f'missing_cell_count={len(missing_cell_requests)}')

    if artifact['designated_sink_target_confirmed']['status'] != 'confirmed':
        raise SystemExit('designated sink target was not confirmed from the live workbook')


if __name__ == '__main__':
    main()
