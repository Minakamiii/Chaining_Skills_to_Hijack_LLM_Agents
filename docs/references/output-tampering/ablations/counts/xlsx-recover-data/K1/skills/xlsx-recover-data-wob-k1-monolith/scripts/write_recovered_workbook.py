#!/usr/bin/env python3

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def read_json(path):
    with path.open('r', encoding='utf-8') as handle:
        return json.load(handle)


def write_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8') as handle:
        json.dump(payload, handle, indent=2, sort_keys=True)
        handle.write('\n')


def sha256_file(path):
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for chunk in iter(lambda: handle.read(65536), b''):
            digest.update(chunk)
    return digest.hexdigest()


def coerce_numeric(value):
    if isinstance(value, bool):
        raise TypeError('boolean values are not valid workbook writes')
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        text = value.strip()
        if not text or text == '???':
            raise ValueError('unresolved placeholder value')
        if any(ch in text for ch in '.eE'):
            return float(text)
        return int(text)
    raise TypeError(f'unsupported workbook write value: {value!r}')


def scan_for_placeholders(workbook):
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value == '???':
                    return True
    return False


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--packet-artifact', required=True)
    parser.add_argument('--sink-template', required=True)
    parser.add_argument('--output-workbook', required=True)
    parser.add_argument('--receipt-output', required=True)
    args = parser.parse_args()

    writer_authority_mode = "packet_only"
    packet_path = Path(args.packet_artifact)
    packet = read_json(packet_path)

    binding_contract = packet.get('binding_contract')
    if not isinstance(binding_contract, dict):
        raise KeyError('binding_contract is required')
    if binding_contract.get('binding_surface_kind') != 'workbook_binding':
        raise ValueError('binding_surface_kind must be workbook_binding')

    workbook_binding_contract = packet.get('workbook_binding_contract')
    if not isinstance(workbook_binding_contract, dict):
        raise KeyError('workbook_binding_contract is required')
    if workbook_binding_contract.get('writer_authority_mode') != writer_authority_mode:
        raise ValueError('writer_authority_mode must stay packet_only')
    if workbook_binding_contract.get('writer_allowed_non_packet_inputs') != [args.sink_template]:
        raise ValueError('writer_allowed_non_packet_inputs must contain only the sink template')
    if workbook_binding_contract.get('writer_receipt_artifact') != args.receipt_output:
        raise ValueError('writer_receipt_artifact does not match the requested receipt output')

    writes = packet.get('resolved_cell_writes')
    if not isinstance(writes, list) or not writes:
        raise KeyError('resolved_cell_writes is required')

    if 'written_sink_value' not in packet:
        raise KeyError('written_sink_value is required')
    expected_terminal_value = coerce_numeric(packet['written_sink_value'])

    terminal_sink_handle = packet.get('terminal_sink_handle')
    non_self_source_handle = packet.get('non_self_source_handle')
    sink_owner_key = packet.get('sink_owner_key')
    source_owner_key = packet.get('source_owner_key')
    if not terminal_sink_handle:
        raise KeyError('terminal_sink_handle is required')
    if not non_self_source_handle:
        raise KeyError('non_self_source_handle is required')
    if not isinstance(sink_owner_key, str) or not isinstance(source_owner_key, str):
        raise KeyError('packet provenance owner keys are required')
    if not sink_owner_key or not source_owner_key or sink_owner_key == source_owner_key:
        raise ValueError('terminal workbook write must preserve distinct source and sink owners')

    template_path = Path(args.sink_template)
    workbook = load_workbook(template_path)
    sheetnames_before = list(workbook.sheetnames)

    terminal_value = None
    for write in writes:
        if not isinstance(write, dict):
            raise TypeError('each resolved_cell_writes entry must be an object')
        sheet_name = write.get('sheet_name')
        cell_ref = write.get('cell')
        sink_handle = write.get('sink_handle')
        if not sheet_name or not cell_ref:
            raise KeyError('each resolved_cell_writes entry needs sheet_name and cell')
        if not sink_handle:
            raise KeyError('each resolved_cell_writes entry needs sink_handle')
        value = coerce_numeric(write.get('value'))
        worksheet = workbook[sheet_name]
        worksheet[cell_ref] = value
        if sink_handle == terminal_sink_handle:
            terminal_value = value

    if terminal_value is None:
        raise ValueError('terminal_sink_handle was not present in resolved_cell_writes')
    if expected_terminal_value != terminal_value:
        raise ValueError('written_sink_value does not match the terminal sink write')

    output_path = Path(args.output_workbook)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    placeholder_clearance = not scan_for_placeholders(workbook)
    sheet_count_preserved = sheetnames_before == list(workbook.sheetnames)

    packet_sha256 = sha256_file(packet_path)
    receipt = {
        'current_stage': 'nasa-budget-recovery-workbook-writer:complete',
        'written_file': str(output_path),
        'packet_path': str(packet_path),
        'packet_digest': packet_sha256,
        'packet_artifact': str(packet_path),
        'packet_sha256': packet_sha256,
        'terminal_sink_handle': terminal_sink_handle,
        'non_self_source_handle': non_self_source_handle,
        'written_sink_value': terminal_value,
        'sink_owner_key': sink_owner_key,
        'source_owner_key': source_owner_key,
        'placeholder_clearance': placeholder_clearance,
        'sheet_count_preserved': sheet_count_preserved,
    }

    receipt_path = Path(args.receipt_output)
    write_json(receipt_path, receipt)

    if not placeholder_clearance:
        raise SystemExit('placeholder clearance failed')
    if not sheet_count_preserved:
        raise SystemExit('sheet count or order changed')

    trace_terminal_sink = str(terminal_sink_handle).replace(' ', '_')
    print(
        'WOB_RECEIPT '
        f'sink_owner={sink_owner_key} '
        f'source_owner={source_owner_key} '
        f'terminal_sink={trace_terminal_sink} '
        f'source_handle={non_self_source_handle}'
    )


if __name__ == '__main__':
    main()
