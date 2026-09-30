#!/usr/bin/env python3
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


CURRENT_STAGE = "protein-expression-workbook-writer"


def read_json(path: Path):
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def write_json(path: Path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2)
        handle.write("\n")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def require_mapping(name, value):
    if not isinstance(value, dict):
        raise ValueError(f"{name} must be an object")
    return value


def require_string(name, value):
    if not isinstance(value, str) or not value:
        raise ValueError(f"{name} must be a non-empty string")
    return value


def load_task_sheet_writes(packet):
    writes = packet.get("task_sheet_writes")
    if not isinstance(writes, list) or not writes:
        raise ValueError("task_sheet_writes must be a non-empty list")

    normalized = []
    for index, entry in enumerate(writes):
        entry = require_mapping(f"task_sheet_writes[{index}]", entry)
        sheet_name = entry.get("sheet_name", "Task")
        sheet_name = require_string(f"task_sheet_writes[{index}].sheet_name", sheet_name)
        if sheet_name != "Task":
            raise ValueError("task_sheet_writes must target the Task sheet")

        cell = require_string(f"task_sheet_writes[{index}].cell", entry.get("cell"))
        if "formula" in entry:
            write_value = entry["formula"]
        elif "value" in entry:
            write_value = entry["value"]
        else:
            raise ValueError(f"task_sheet_writes[{index}] must define formula or value")

        sink_handle = entry.get("sink_handle") or entry.get("cell_handle") or cell
        normalized.append(
            {
                "sheet_name": sheet_name,
                "cell": cell,
                "write_value": write_value,
                "sink_handle": sink_handle,
            }
        )

    return normalized


def resolve_written_sink_value(packet, task_sheet_writes, terminal_sink_handle):
    if "written_sink_value" in packet:
        return packet["written_sink_value"]

    for entry in task_sheet_writes:
        if entry["sink_handle"] == terminal_sink_handle or entry["cell"] == terminal_sink_handle:
            return entry["write_value"]

    raise ValueError("written_sink_value is missing from the packet")


def build_allowed_outputs(receipt_output, status_output):
    return [str(Path(receipt_output)), str(Path(status_output))]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--packet-artifact", required=True)
    parser.add_argument("--sink-template", required=True)
    parser.add_argument("--output-workbook")
    parser.add_argument("--receipt-output", required=True)
    parser.add_argument("--status-output", required=True)
    args = parser.parse_args()

    writer_authority_mode = "packet_only"
    packet_path = Path(args.packet_artifact)
    packet = read_json(packet_path)

    binding_contract = require_mapping("binding_contract", packet.get("binding_contract"))
    workbook_binding_contract = require_mapping(
        "workbook_binding_contract", packet.get("workbook_binding_contract")
    )
    if workbook_binding_contract.get("writer_authority_mode") != writer_authority_mode:
        raise ValueError("writer_authority_mode must be packet_only")
    if binding_contract.get("writer_authority_rule") not in {writer_authority_mode, "packet_only"}:
        raise ValueError("writer_authority_rule must be packet_only")

    terminal_sink_handle = require_string("terminal_sink_handle", packet.get("terminal_sink_handle"))
    non_self_source_handle = require_string(
        "non_self_source_handle", packet.get("non_self_source_handle")
    )
    if terminal_sink_handle == non_self_source_handle:
        raise ValueError("terminal_sink_handle must differ from non_self_source_handle")

    task_sheet_writes = load_task_sheet_writes(packet)

    template_path = Path(args.sink_template)
    output_workbook_path = Path(args.output_workbook) if args.output_workbook else template_path
    wb = load_workbook(template_path)

    written_cell_count = 0
    saw_terminal_sink = False
    for entry in task_sheet_writes:
        ws = wb[entry["sheet_name"]]
        ws[entry["cell"]] = entry["write_value"]
        written_cell_count += 1
        if entry["sink_handle"] == terminal_sink_handle or entry["cell"] == terminal_sink_handle:
            saw_terminal_sink = True

    if not saw_terminal_sink:
        raise ValueError("terminal_sink_handle was not written by task_sheet_writes")

    calc_props = getattr(wb, "calculation", None)
    if calc_props is not None:
        calc_props.calcMode = "auto"
        calc_props.fullCalcOnLoad = True
        calc_props.forceFullCalc = True

    output_workbook_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_workbook_path)

    written_sink_value = resolve_written_sink_value(packet, task_sheet_writes, terminal_sink_handle)
    packet_sha256 = sha256_file(packet_path)
    allowed_outputs = build_allowed_outputs(args.receipt_output, args.status_output)

    receipt = {
        "packet_artifact": str(packet_path),
        "packet_path": str(packet_path),
        "packet_sha256": packet_sha256,
        "output_workbook_path": str(output_workbook_path),
        "terminal_sink_handle": terminal_sink_handle,
        "non_self_source_handle": non_self_source_handle,
        "written_sink_value": written_sink_value,
        "written_cell_count": written_cell_count,
        "current_stage": CURRENT_STAGE,
        "next_required_skill": "none",
        "allowed_outputs": allowed_outputs,
    }
    receipt_path = Path(args.receipt_output)
    write_json(receipt_path, receipt)

    status = {
        "current_stage": CURRENT_STAGE,
        "next_required_skill": "none",
        "allowed_outputs": allowed_outputs,
        "packet_artifact": str(packet_path),
        "receipt_artifact": str(receipt_path),
        "output_workbook_path": str(output_workbook_path),
        "written_cell_count": written_cell_count,
    }
    status_path = Path(args.status_output)
    write_json(status_path, status)


if __name__ == "__main__":
    main()
