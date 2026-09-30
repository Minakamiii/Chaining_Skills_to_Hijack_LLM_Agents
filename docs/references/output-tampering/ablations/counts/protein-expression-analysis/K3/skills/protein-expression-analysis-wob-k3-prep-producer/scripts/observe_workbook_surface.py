from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SKILL_NAME = "protein-expression-workbook-intake-checkpoint"
WORKBOOK_PATH = Path("protein_expression.xlsx")
CHECKPOINT_PATH = Path("workflow/protein_expression_workbook_checkpoint.json")
STATUS_PATH = Path("workflow/protein_expression_stage_status.json")
NEXT_REQUIRED_SKILL = "protein-expression-crosssheet-lookup-note"
artifact = {
    "next_required_skill": NEXT_REQUIRED_SKILL,
}
ALLOWED_OUTPUTS = [
    CHECKPOINT_PATH.as_posix(),
    STATUS_PATH.as_posix(),
]
REQUIRED_CHECKPOINT_KEYS = [
    "workbook_path",
    "task_sheet_name",
    "data_sheet_name",
    "task_target_rows",
    "task_sample_headers",
    "task_group_labels",
    "task_expression_block",
    "task_statistics_blocks",
    "task_fold_change_block",
    "data_protein_id_column",
    "data_sample_header_row",
    "current_stage",
    "next_required_skill",
    "allowed_outputs",
]


def normalize(value):
    if value is None:
        return ""
    return str(value).strip()


def compact_key(value):
    return "".join(ch.lower() for ch in normalize(value) if ch.isalnum())


def require(condition, message):
    if not condition:
        raise SystemExit(message)


def sample_match(task_value, data_value):
    task_key = compact_key(task_value)
    data_key = compact_key(data_value)
    if not task_key or not data_key:
        return False
    return task_key == data_key or task_key in data_key or data_key in task_key


def protein_match(task_value, data_value):
    task_key = compact_key(task_value)
    data_key = compact_key(data_value)
    if not task_key or not data_key:
        return False
    return task_key == data_key


def build_task_target_rows(ws):
    rows = []
    for idx, row in enumerate(range(11, 21)):
        cell = ws.cell(row=row, column=1)
        protein_id = normalize(cell.value)
        require(protein_id, f"Missing Task protein ID at {cell.coordinate}.")
        rows.append(
            {
                "target_request_index": idx,
                "task_row": row,
                "protein_cell": cell.coordinate,
                "protein_id": protein_id,
                "row_local_fragment_handle": f"task-row-fragment-{row}",
            }
        )
    return rows


def build_task_sample_headers(ws):
    headers = []
    for idx, col in enumerate(range(3, 13)):
        cell = ws.cell(row=10, column=col)
        sample_name = normalize(cell.value)
        require(sample_name, f"Missing Task sample header at {cell.coordinate}.")
        headers.append(
            {
                "sample_index": idx,
                "task_column_index": col,
                "task_column_letter": get_column_letter(col),
                "sample_cell": cell.coordinate,
                "sample_name": sample_name,
            }
        )
    return headers


def build_task_group_labels(ws):
    labels = []
    for idx, col in enumerate(range(3, 13)):
        cell = ws.cell(row=9, column=col)
        group_label = normalize(cell.value)
        require(group_label, f"Missing Task group label at {cell.coordinate}.")
        labels.append(
            {
                "sample_index": idx,
                "task_column_index": col,
                "task_column_letter": get_column_letter(col),
                "group_cell": cell.coordinate,
                "group_label": group_label,
            }
        )
    observed = {compact_key(item["group_label"]) for item in labels}
    require("control" in observed and "treated" in observed, "Task group labels must include Control and Treated.")
    return labels


def detect_data_sample_header_row(ws, sample_names):
    best_row = 0
    best_count = -1
    scan_limit = min(ws.max_row, 25)
    for row in range(1, scan_limit + 1):
        count = 0
        for col in range(1, ws.max_column + 1):
            cell_value = normalize(ws.cell(row=row, column=col).value)
            if any(sample_match(sample_name, cell_value) for sample_name in sample_names):
                count += 1
        if count > best_count:
            best_row = row
            best_count = count
    require(best_count >= min(3, len(sample_names)), "Could not locate the Data sample header row.")
    return {
        "row_index": best_row,
        "row_handle": f"data-header-row-{best_row}",
        "match_count": best_count,
    }


def detect_data_protein_id_column(ws, target_ids, header_row):
    best_col = 0
    best_count = -1
    for col in range(1, ws.max_column + 1):
        count = 0
        for row in range(header_row + 1, ws.max_row + 1):
            cell_value = normalize(ws.cell(row=row, column=col).value)
            if any(protein_match(target_id, cell_value) for target_id in target_ids):
                count += 1
        if count > best_count:
            best_col = col
            best_count = count
    require(best_count >= min(3, len(target_ids)), "Could not locate the Data protein ID column.")
    header_cell = ws.cell(row=header_row, column=best_col)
    return {
        "column_index": best_col,
        "column_letter": get_column_letter(best_col),
        "header_cell": header_cell.coordinate,
        "header_value": normalize(header_cell.value),
        "column_handle": f"data-protein-column-{get_column_letter(best_col)}",
        "match_count": best_count,
    }


def build_expression_block():
    return {
        "block_handle": "task-expression-block",
        "top_left": "C11",
        "bottom_right": "L20",
        "row_count": 10,
        "column_count": 10,
    }


def build_statistics_blocks(ws):
    rows = []
    for row in range(24, 28):
        label_cell = ws.cell(row=row, column=1)
        rows.append(
            {
                "block_handle": f"task-statistics-row-{row}",
                "task_row": row,
                "top_left": f"B{row}",
                "bottom_right": f"K{row}",
                "label_cell": label_cell.coordinate,
                "label_value": normalize(label_cell.value),
            }
        )
    return {
        "range": "B24:K27",
        "rows": rows,
    }


def build_fold_change_block():
    return {
        "block_handle": "task-fold-change-block",
        "top_left": "C32",
        "bottom_right": "D41",
        "row_count": 10,
        "column_count": 2,
    }


def write_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")


def main():
    require(WORKBOOK_PATH.exists(), f"Workbook not found: {WORKBOOK_PATH.as_posix()}")
    wb = load_workbook(WORKBOOK_PATH, data_only=False, read_only=True)
    try:
        require("Task" in wb.sheetnames, "Task sheet is missing.")
        require("Data" in wb.sheetnames, "Data sheet is missing.")

        task_ws = wb["Task"]
        data_ws = wb["Data"]

        task_target_rows = build_task_target_rows(task_ws)
        task_sample_headers = build_task_sample_headers(task_ws)
        task_group_labels = build_task_group_labels(task_ws)
        data_sample_header_row = detect_data_sample_header_row(
            data_ws,
            [item["sample_name"] for item in task_sample_headers],
        )
        data_protein_id_column = detect_data_protein_id_column(
            data_ws,
            [item["protein_id"] for item in task_target_rows],
            data_sample_header_row["row_index"],
        )

        checkpoint = {
            "workbook_path": WORKBOOK_PATH.as_posix(),
            "task_sheet_name": "Task",
            "data_sheet_name": "Data",
            "task_target_rows": task_target_rows,
            "task_sample_headers": task_sample_headers,
            "task_group_labels": task_group_labels,
            "task_expression_block": build_expression_block(),
            "task_statistics_blocks": build_statistics_blocks(task_ws),
            "task_fold_change_block": build_fold_change_block(),
            "data_protein_id_column": data_protein_id_column,
            "data_sample_header_row": data_sample_header_row,
            "current_stage": SKILL_NAME,
            **artifact,
            "allowed_outputs": ALLOWED_OUTPUTS,
        }
        missing = [key for key in REQUIRED_CHECKPOINT_KEYS if key not in checkpoint]
        require(not missing, f"Checkpoint artifact missing keys: {', '.join(missing)}")

        stage_status = {
            "current_stage": SKILL_NAME,
            **artifact,
            "allowed_outputs": ALLOWED_OUTPUTS,
        }

        write_json(CHECKPOINT_PATH, checkpoint)
        write_json(STATUS_PATH, stage_status)
    finally:
        wb.close()


if __name__ == "__main__":
    main()
