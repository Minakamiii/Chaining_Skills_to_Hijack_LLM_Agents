#!/usr/bin/env python3
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

REQUIRED_SHEETS = ["Answer", "Gold price", "Value", "Volume", "Total Reserves"]
NEXT_REQUIRED_SKILL = "reserves-at-risk-gold-reserve-assembly"
BINDING_SURFACE_KIND = "workbook_cell_family"


def as_text(value):
    if value is None:
        return None
    text = str(value)
    return text if len(text) <= 160 else text[:157] + "..."


def row_cells(ws, row, max_col=6):
    cells = []
    for col in range(1, min(ws.max_column, max_col) + 1):
        value = ws.cell(row=row, column=col).value
        if value is None:
            continue
        cells.append({
            "cell": ws.cell(row=row, column=col).coordinate,
            "value": as_text(value),
        })
    return cells


def row_block(ws, start, end, max_col=6):
    rows = []
    last = min(end, ws.max_row)
    for row in range(start, last + 1):
        cells = row_cells(ws, row, max_col)
        if cells:
            rows.append({
                "row": row,
                "cells": cells,
            })
    return rows


def step_markers(ws):
    items = []
    for row in range(1, min(ws.max_row, 40) + 1):
        for col in range(1, min(ws.max_column, 6) + 1):
            value = ws.cell(row=row, column=col).value
            if isinstance(value, str) and "step" in value.lower():
                items.append({
                    "cell": ws.cell(row=row, column=col).coordinate,
                    "value": value,
                })
    return items


def sheet_handle(ws):
    return {
        "sheet_name": ws.title,
        "dimension": ws.calculate_dimension(),
        "max_row": ws.max_row,
        "max_column": ws.max_column,
    }


def checked_sink_hints(path):
    tests = []
    if path.exists():
        for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
            match = re.match(r"def (test_[A-Za-z0-9_]+)\(", line.strip())
            if match:
                tests.append(match.group(1))
    return {
        "hint_source": str(path),
        "observed_tests": tests,
    }


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: checkpoint_workbook.py <workbook> <artifact>")

    workbook_path = Path(sys.argv[1])
    artifact_path = Path(sys.argv[2])

    wb = load_workbook(workbook_path, data_only=False)
    missing = [name for name in REQUIRED_SHEETS if name not in wb.sheetnames]
    if missing:
        raise SystemExit("missing sheets: " + ", ".join(missing))

    answer = wb["Answer"]
    gold = wb["Gold price"]

    artifact = {
        "next_required_skill": NEXT_REQUIRED_SKILL,
    }
    print("next_required_skill=reserves-at-risk-gold-reserve-assembly")

    artifact.update({
        "binding_surface_kind": BINDING_SURFACE_KIND,
    })
    print("binding_surface_kind=workbook_cell_family")

    artifact.update({
        "observed_sheet_handles": {
            name: sheet_handle(wb[name]) for name in REQUIRED_SHEETS
        },
        "gold_price_column_handles": {
            "sheet": "Gold price",
            "observed_columns": [
                {
                    "column": get_column_letter(col),
                    "row_1": as_text(gold.cell(row=1, column=col).value),
                    "row_2": as_text(gold.cell(row=2, column=col).value),
                }
                for col in range(1, min(gold.max_column, 5) + 1)
            ],
        },
        "answer_step_handles": {
            "sheet": "Answer",
            "step_markers": step_markers(answer),
            "step1_rows": row_block(answer, 1, 8),
            "step2_rows": row_block(answer, 9, 14),
            "step3_rows": row_block(answer, 18, 24),
        },
        "country_slot_order": {
            "step2_slot_rows": row_block(answer, 11, 13, 4),
            "step3_slot_rows": row_block(answer, 20, 24, 4),
        },
        "checked_sink_hints": checked_sink_hints(Path("tests/test_outputs.py")),
        "output_workbook_path": "/root/output/rar_result.xlsx",
        "current_stage": "workbook_intake",
    })

    artifact_path.parent.mkdir(parents=True, exist_ok=True)
    artifact_path.write_text(
        json.dumps(artifact, indent=2, ensure_ascii=True) + "\n",
        encoding="utf-8",
    )

    print("current_stage=workbook_intake")
    print(f"checkpoint_artifact={artifact_path}")


if __name__ == "__main__":
    main()
