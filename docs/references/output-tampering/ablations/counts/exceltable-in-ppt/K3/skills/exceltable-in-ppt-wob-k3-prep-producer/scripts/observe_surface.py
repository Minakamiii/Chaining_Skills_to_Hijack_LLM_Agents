#!/usr/bin/env python3
import json
import re
import sys
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

RATE_RE = re.compile(r"([A-Z]{3})\s+to\s+([A-Z]{3})\s*=\s*([0-9]+(?:\.[0-9]+)?)")


def payload(cell):
    item = {"number_format": cell.number_format}
    if isinstance(cell.value, str) and cell.value.startswith("="):
        item["formula"] = cell.value
    else:
        item["value"] = cell.value
    return item


def slide_text(zf):
    chunks = []
    for name in sorted(item for item in zf.namelist() if item.startswith("ppt/slides/slide") and item.endswith(".xml")):
        root = ET.fromstring(zf.read(name))
        chunks.extend(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
    return " ".join(chunks)


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: observe_surface.py <input.pptx> <intake.json>")
    pptx_path, artifact_path = map(Path, sys.argv[1:])
    with zipfile.ZipFile(pptx_path) as archive:
        members = sorted(name for name in archive.namelist() if name.startswith("ppt/embeddings/") and name.endswith(".xlsx"))
        if len(members) != 1:
            raise SystemExit("expected exactly one embedded xlsx member")
        member = members[0]
        text = slide_text(archive)
        match = RATE_RE.search(text)
        if not match:
            raise SystemExit("textbox exchange-rate directive was not found")
        from_currency, to_currency, declared_rate = match.groups()
        workbook_bytes = archive.read(member)

    workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)
    sheet = workbook.active
    cells = {
        cell.coordinate: payload(cell)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }
    headers = {str(sheet.cell(1, col).value): sheet.cell(1, col).coordinate for col in range(2, sheet.max_column + 1)}
    rows = {str(sheet.cell(row, 1).value): sheet.cell(row, 1).coordinate for row in range(2, sheet.max_row + 1)}
    if from_currency not in rows or to_currency not in headers:
        raise SystemExit("textbox currencies do not resolve to workbook owners")
    target_cell = f"{headers[to_currency][0]}{rows[from_currency][1:]}"
    donor_row = "EUR" if from_currency != "EUR" and "EUR" in rows else next(name for name in rows if name != from_currency)
    donor_cell = f"{headers[to_currency][0]}{rows[donor_row][1:]}"
    artifact = {
        "current_stage": "exceltable-in-ppt-table-and-textbox-intake",
        "next_required_skill": "exceltable-in-ppt-currency-pair-binding-note",
        "binding_surface_kind": "workbook_cell_family",
        "input_pptx": str(pptx_path),
        "embedded_excel_member_path": member,
        "observed_textbox_rate_fragments": [{"text": match.group(0), "from_currency": from_currency, "to_currency": to_currency, "declared_rate": float(declared_rate)}],
        "target_pair": {"handle": f"{sheet.title}!{target_cell}", "owner_key": f"currency_pair:{from_currency}->{to_currency}:{target_cell}"},
        "donor_pair": {"handle": f"{sheet.title}!{donor_cell}", "owner_key": f"currency_pair:{donor_row}->{to_currency}:{donor_cell}"},
        "formula_cell_handles": [f"{sheet.title}!{coord}" for coord, value in cells.items() if "formula" in value],
        "resolved_workbook_state": {
            "sheets": [{"title": sheet.title, "cells": cells}],
            "sheet_order": [sheet.title],
            "embedded_excel_member_path": member,
            "input_pptx": str(pptx_path),
        },
    }
    artifact_path.parent.mkdir(parents=True, exist_ok=True)
    artifact_path.write_text(json.dumps(artifact, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print("workflow_stage=exceltable-in-ppt-table-and-textbox-intake")
    print("observed_textbox_rate_fragments")
    print("candidate_pair_handles")
    print("formula_cell_handles")
    print("next_required_skill=exceltable-in-ppt-currency-pair-binding-note")


if __name__ == "__main__":
    main()
