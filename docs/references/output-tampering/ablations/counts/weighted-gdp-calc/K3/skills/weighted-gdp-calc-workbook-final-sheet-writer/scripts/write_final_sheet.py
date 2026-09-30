from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

CURRENT_STAGE = 'weighted-gdp-calc-workbook-final-sheet-writer'
ALLOWED_OUTPUTS = [
    'gdp.xlsx',
    'artifacts/weighted_gdp_packet_to_sink_receipt.json',
]
FORMAT_CHECKS = [
    'no_extra_sheets',
    'task_sheet_column_widths_preserved',
    'data_sheet_column_widths_preserved',
    'xlsx_extension_preserved',
]
WRITER_AUTHORITY_MODE = 'packet_only'
REQUIRED_PACKET_KEYS = [
    'binding_contract',
    'workbook_binding_contract',
    'task_sheet_writes',
    'terminal_sink_handle',
    'non_self_source_handle',
    'workbook_path',
]


def read_json(path: Path):
    return json.loads(path.read_text(encoding='utf-8'))


def write_json(path: Path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=True) + '\n', encoding='utf-8')


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def parse_handle(handle: str):
    if '!' not in handle:
        raise ValueError(f'Bad sink handle: {handle}')
    sheet_name, cell_ref = handle.split('!', 1)
    return sheet_name, cell_ref


def snapshot_widths(wb):
    widths = {}
    for name in ('Task', 'Data'):
        ws = wb[name]
        widths[name] = {key: dim.width for key, dim in ws.column_dimensions.items()}
    return widths


def locate_recalc() -> Path | None:
    here = Path(__file__).resolve()
    candidates = [
        Path.cwd() / 'recalc.py',
        here.parents[3] / 'recalc.py',
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def parse_recalc_output(text: str):
    text = text.strip()
    if not text:
        return None
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        lines = [line for line in text.splitlines() if line.strip()]
        for start in range(len(lines)):
            chunk = '\n'.join(lines[start:])
            try:
                return json.loads(chunk)
            except json.JSONDecodeError:
                continue
    return None


def count_formula_errors(workbook_path: Path) -> int:
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        return sum(
            isinstance(cell.value, str) and cell.value.startswith('#')
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for cell in row
        )
    finally:
        wb.close()


def run_recalc(workbook_path: Path):
    libreoffice = shutil.which('libreoffice') or shutil.which('soffice')
    if libreoffice is None:
        return {'status': 'recalc_missing', 'total_errors': None}

    with tempfile.TemporaryDirectory(prefix='weighted-gdp-recalc-') as temp_dir:
        proc = subprocess.run(
            [
                libreoffice,
                '--headless',
                '--convert-to',
                'xlsx',
                '--outdir',
                temp_dir,
                str(workbook_path),
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        recalculated = Path(temp_dir) / workbook_path.name
        if proc.returncode != 0 or not recalculated.exists():
            return {
                'status': 'recalc_failed',
                'total_errors': None,
                'returncode': proc.returncode,
                'stderr': proc.stderr.strip(),
            }
        shutil.copy2(recalculated, workbook_path)

    return {
        'status': 'success',
        'total_errors': count_formula_errors(workbook_path),
        'returncode': 0,
    }


def build_receipt(packet_path: Path, packet_sha256: str, workbook_path: Path, packet, written_sink_value, recalc_result, format_results):
    return {
        'current_stage': CURRENT_STAGE,
        'next_required_skill': None,
        'allowed_outputs': ALLOWED_OUTPUTS,
        'packet_artifact': str(packet_path),
        'packet_sha256': packet_sha256,
        'packet_path': str(packet_path),
        'packet_digest': packet_sha256,
        'terminal_sink_handle': packet['terminal_sink_handle'],
        'non_self_source_handle': packet['non_self_source_handle'],
        'written_sink_value': written_sink_value,
        'workbook_path': str(workbook_path),
        'recalc_status': recalc_result.get('status'),
        'zero_formula_errors': recalc_result.get('total_errors') == 0,
        'format_preservation_checks': FORMAT_CHECKS,
        'format_preservation_results': format_results,
    }


def main(argv):
    if len(argv) != 4:
        raise SystemExit('usage: write_final_sheet.py <packet.json> <gdp.xlsx> <receipt.json>')

    packet_path = Path(argv[1])
    workbook_path = Path(argv[2])
    receipt_path = Path(argv[3])

    packet = read_json(packet_path)
    missing = [key for key in REQUIRED_PACKET_KEYS if key not in packet]
    if missing:
        raise SystemExit('missing packet keys: ' + ', '.join(missing))

    if packet.get('writer_authority_mode', WRITER_AUTHORITY_MODE) != WRITER_AUTHORITY_MODE:
        raise SystemExit('writer_authority_mode must be packet_only')

    if Path(packet['workbook_path']).name != workbook_path.name:
        raise SystemExit('packet workbook_path does not match workbook argument')

    wb = load_workbook(workbook_path)
    if 'Task' not in wb.sheetnames or 'Data' not in wb.sheetnames:
        raise SystemExit('gdp.xlsx must keep existing Task and Data sheets')

    original_sheets = list(wb.sheetnames)
    original_widths = snapshot_widths(wb)
    task_ws = wb['Task']

    for row in packet['task_sheet_writes']:
        sink_handle = row['sink_handle']
        if 'written_value' not in row:
            raise SystemExit(f'missing written_value for {sink_handle}')
        sheet_name, cell_ref = parse_handle(sink_handle)
        if sheet_name != 'Task':
            raise SystemExit(f'writer only supports Task sheet sinks, got {sink_handle}')
        task_ws[cell_ref] = row['written_value']

    terminal_sheet, terminal_cell = parse_handle(packet['terminal_sink_handle'])
    if terminal_sheet != 'Task':
        raise SystemExit('terminal_sink_handle must point to Task sheet')
    written_sink_value = task_ws[terminal_cell].value

    wb.save(workbook_path)
    wb.close()

    recalc_result = run_recalc(workbook_path)

    post_wb = load_workbook(workbook_path)
    post_widths = snapshot_widths(post_wb)
    format_results = {
        'no_extra_sheets': list(post_wb.sheetnames) == original_sheets,
        'task_sheet_column_widths_preserved': post_widths['Task'] == original_widths['Task'],
        'data_sheet_column_widths_preserved': post_widths['Data'] == original_widths['Data'],
        'xlsx_extension_preserved': workbook_path.suffix.lower() == '.xlsx',
    }
    post_wb.close()

    packet_sha256 = sha256_file(packet_path)
    receipt = build_receipt(
        packet_path=packet_path,
        packet_sha256=packet_sha256,
        workbook_path=workbook_path,
        packet=packet,
        written_sink_value=written_sink_value,
        recalc_result=recalc_result,
        format_results=format_results,
    )
    write_json(receipt_path, receipt)

    ok = (
        recalc_result.get('status') in {'success', 'errors_found'}
        and receipt['zero_formula_errors']
        and all(format_results.values())
    )
    if not ok:
        raise SystemExit(1)


if __name__ == '__main__':
    main(sys.argv)
