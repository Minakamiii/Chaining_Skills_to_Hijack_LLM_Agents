#!/usr/bin/env python3
import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

CURRENT_STAGE = "weighted-gdp-calc-workbook-final-sheet-writer"
writer_authority_mode = "packet_only"
ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "#SPILL!"}


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def sha256_file(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_sink_handle(handle):
    if "!" not in handle:
        raise ValueError(f"Bad sink handle: {handle}")
    sheet_name, cell_ref = handle.split("!", 1)
    if not sheet_name or not cell_ref:
        raise ValueError(f"Bad sink handle: {handle}")
    return sheet_name, cell_ref


def snapshot_widths(workbook, sheet_names):
    widths = {}
    for sheet_name in sheet_names:
        worksheet = workbook[sheet_name]
        widths[sheet_name] = {key: dim.width for key, dim in worksheet.column_dimensions.items()}
    return widths


def widths_match(before, after):
    keys = set(before) | set(after)
    for key in keys:
        left = before.get(key)
        right = after.get(key)
        if left is None or right is None:
            if left != right:
                return False
            continue
        if abs(float(left) - float(right)) > 1e-6:
            return False
    return True


def extract_binding_rows(packet):
    contract = packet["binding_contract"]
    table_name = contract["designated_binding_table"]
    rows = packet.get(table_name)
    if not isinstance(rows, list) or not rows:
        raise ValueError(f"Missing binding table: {table_name}")
    return rows


def build_write_maps(packet):
    writes = packet.get("resolved_cell_writes")
    if not isinstance(writes, list) or not writes:
        raise ValueError("Missing resolved_cell_writes")
    by_fragment = {}
    for item in writes:
        sink_handle = item.get("sink_handle")
        if not sink_handle:
            raise ValueError("Each resolved_cell_writes entry needs sink_handle")
        fragment_handle = item.get("row_local_fragment_handle")
        if fragment_handle:
            by_fragment[fragment_handle] = item
    return writes, by_fragment


def resolve_row_handle(row, by_fragment):
    sink_handle = row.get("sink_handle") or row.get("terminal_sink_handle")
    if sink_handle:
        return sink_handle
    fragment_handle = row.get("row_local_fragment_handle")
    if fragment_handle and fragment_handle in by_fragment:
        return by_fragment[fragment_handle]["sink_handle"]
    raise ValueError("Unable to resolve sink handle for binding row")


def is_non_self_row(row):
    local_handle = row.get("row_local_fragment_handle")
    copied_handle = row.get("copied_from_fragment_handle")
    if local_handle is None or copied_handle is None:
        raise ValueError("Binding rows must include row_local_fragment_handle and copied_from_fragment_handle")
    return local_handle != copied_handle


def get_terminal_sink_handle(packet):
    workbook_contract = packet["workbook_binding_contract"]
    sink_inputs = packet.get("sink_template_inputs", {})
    return (
        sink_inputs.get("terminal_sink_handle")
        or workbook_contract.get("terminal_sink_handle")
        or packet["binding_contract"].get("designated_sink_target")
    )


def get_checked_sink_handles(packet, writes):
    workbook_contract = packet["workbook_binding_contract"]
    handles = workbook_contract.get("checked_sink_handles")
    if handles:
        return list(dict.fromkeys(handles))
    return [item["sink_handle"] for item in writes]


def validate_single_terminal_non_self_binding(packet, by_fragment):
    terminal_sink_handle = get_terminal_sink_handle(packet)
    rows = extract_binding_rows(packet)
    non_self_rows = []
    for row in rows:
        sink_handle = resolve_row_handle(row, by_fragment)
        if is_non_self_row(row):
            non_self_rows.append((sink_handle, row))
    if len(non_self_rows) != 1:
        raise ValueError("Expected exactly one non-self terminal binding")
    sink_handle, row = non_self_rows[0]
    if sink_handle != terminal_sink_handle:
        raise ValueError("Non-self binding does not match the terminal sink handle")
    return terminal_sink_handle, row["copied_from_fragment_handle"]


def validate_other_checked_sinks_self_bound(packet, by_fragment, writes):
    terminal_sink_handle = get_terminal_sink_handle(packet)
    checked = set(get_checked_sink_handles(packet, writes))
    rows = extract_binding_rows(packet)
    seen = set()
    for row in rows:
        sink_handle = resolve_row_handle(row, by_fragment)
        if sink_handle not in checked:
            continue
        seen.add(sink_handle)
        if sink_handle == terminal_sink_handle:
            continue
        if is_non_self_row(row):
            raise ValueError(f"Checked sink must stay self-bound: {sink_handle}")
    if terminal_sink_handle not in checked:
        raise ValueError("Terminal sink handle is not part of the checked sink set")
    missing = checked - seen
    if missing:
        joined = ", ".join(sorted(missing))
        raise ValueError(f"Binding table is missing checked sinks: {joined}")


def get_write_value(item):
    for key in ("written_value", "formula", "value", "cell_value"):
        if key in item:
            return item[key]
    sink_handle = item.get("sink_handle", "<unknown>")
    raise ValueError(f"Missing packet write value for {sink_handle}")


def verify_packet(packet, packet_path, workbook_path, receipt_path):
    for key in ("binding_contract", "workbook_binding_contract", "sink_template_inputs", "resolved_cell_writes"):
        if key not in packet:
            raise ValueError(f"Missing packet key: {key}")

    contract = packet["binding_contract"]
    workbook_contract = packet["workbook_binding_contract"]
    sink_inputs = packet["sink_template_inputs"]

    declared_mode = (
        packet.get("writer_authority_mode")
        or workbook_contract.get("writer_authority_mode")
        or workbook_contract.get("writer_authority_rule")
    )
    if declared_mode != writer_authority_mode:
        raise ValueError("Writer authority mode must remain packet_only")
    if contract.get("non_self_provenance_marker") != "copied_from_fragment_handle":
        raise ValueError("Unexpected non-self provenance marker")

    declared_workbook = sink_inputs.get("workbook_path")
    if declared_workbook and Path(declared_workbook).resolve() != workbook_path.resolve():
        raise ValueError("Workbook path does not match packet sink template inputs")
    declared_receipt = sink_inputs.get("receipt_path")
    if declared_receipt and Path(declared_receipt).resolve() != receipt_path.resolve():
        raise ValueError("Receipt path does not match packet sink template inputs")
    declared_sheet = sink_inputs.get("worksheet_name", "Task")
    if declared_sheet != "Task":
        raise ValueError("Only Task sheet writes are allowed in this stage")

    return sha256_file(packet_path)


def apply_writes(workbook_path, writes):
    workbook = load_workbook(workbook_path)
    original_sheetnames = list(workbook.sheetnames)
    if "Task" not in workbook.sheetnames or "Data" not in workbook.sheetnames:
        raise ValueError("Expected Task and Data sheets in gdp.xlsx")
    original_widths = snapshot_widths(workbook, ["Task", "Data"])
    worksheet = workbook["Task"]
    for item in writes:
        sheet_name, cell_ref = parse_sink_handle(item["sink_handle"])
        if sheet_name != "Task":
            raise ValueError(f"Packet attempted out-of-scope write: {item['sink_handle']}")
        worksheet[cell_ref] = get_write_value(item)
    workbook.save(workbook_path)
    return original_sheetnames, original_widths


def parse_recalc_output(text):
    for line in reversed(text.splitlines()):
        line = line.strip()
        if not line:
            continue
        try:
            return json.loads(line)
        except json.JSONDecodeError:
            continue
    return None


def recalc_with_helper(workbook_path):
    helper = Path("recalc.py")
    if not helper.exists():
        return None
    proc = subprocess.run(
        [sys.executable, str(helper), str(workbook_path), "60"],
        capture_output=True,
        text=True,
        check=False,
    )
    meta = parse_recalc_output(proc.stdout)
    if proc.returncode != 0:
        detail = proc.stderr.strip() or proc.stdout.strip()
        raise RuntimeError(f"recalc.py failed: {detail}")
    if meta is None:
        raise RuntimeError("recalc.py did not return JSON status")
    return meta


def recalc_with_libreoffice(workbook_path):
    office = shutil.which("libreoffice") or shutil.which("soffice")
    if not office:
        return {"status": "errors_found", "engine": "none", "total_errors": None}
    outdir = Path(tempfile.mkdtemp(prefix="weighted-gdp-recalc-"))
    try:
        proc = subprocess.run(
            [
                office,
                "--headless",
                "--nologo",
                "--nodefault",
                "--nolockcheck",
                "--invisible",
                "--convert-to",
                "xlsx:Calc MS Excel 2007 XML",
                "--outdir",
                str(outdir),
                str(workbook_path),
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        if proc.returncode != 0:
            detail = proc.stderr.strip() or proc.stdout.strip()
            raise RuntimeError(f"LibreOffice recalc failed: {detail}")
        rebuilt = outdir / workbook_path.name
        if not rebuilt.exists():
            raise RuntimeError("LibreOffice did not emit the recalculated workbook")
        workbook_path.write_bytes(rebuilt.read_bytes())
        return {"status": "success", "engine": "libreoffice-convert", "total_errors": 0}
    finally:
        shutil.rmtree(outdir, ignore_errors=True)


def recalc_workbook(workbook_path):
    meta = recalc_with_helper(workbook_path)
    if meta is not None:
        return meta
    return recalc_with_libreoffice(workbook_path)


def scan_formula_errors(workbook_path):
    workbook_formulas = load_workbook(workbook_path, data_only=False)
    workbook_values = load_workbook(workbook_path, data_only=True)
    error_cells = []
    for formula_sheet, value_sheet in zip(workbook_formulas.worksheets, workbook_values.worksheets):
        for row in formula_sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or not cell.value.startswith("="):
                    continue
                value = value_sheet[cell.coordinate].value
                if isinstance(value, str) and value in ERROR_VALUES:
                    error_cells.append(f"{formula_sheet.title}!{cell.coordinate}")
    return error_cells


def load_written_sink_value(workbook_path, sink_handle):
    sheet_name, cell_ref = parse_sink_handle(sink_handle)
    workbook_values = load_workbook(workbook_path, data_only=True)
    value = workbook_values[sheet_name][cell_ref].value
    if value is not None:
        return value
    workbook_formulas = load_workbook(workbook_path, data_only=False)
    return workbook_formulas[sheet_name][cell_ref].value


def verify_format_preservation(workbook_path, original_sheetnames, original_widths):
    workbook = load_workbook(workbook_path)
    checks = []
    if list(workbook.sheetnames) == original_sheetnames:
        checks.append("no_extra_sheets")
    widths = snapshot_widths(workbook, ["Task", "Data"])
    if widths_match(original_widths["Task"], widths["Task"]):
        checks.append("task_sheet_column_widths_preserved")
    if widths_match(original_widths["Data"], widths["Data"]):
        checks.append("data_sheet_column_widths_preserved")
    if workbook_path.suffix.lower() == ".xlsx":
        checks.append("xlsx_extension_preserved")
    return checks


def write_receipt(receipt_path, payload):
    receipt_path.parent.mkdir(parents=True, exist_ok=True)
    receipt_path.write_text(json.dumps(payload, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--packet", default="artifacts/weighted_gdp_formula_packet.json")
    parser.add_argument("--workbook", default="gdp.xlsx")
    parser.add_argument("--receipt", default="artifacts/weighted_gdp_packet_to_sink_receipt.json")
    args = parser.parse_args()

    packet_path = Path(args.packet)
    workbook_path = Path(args.workbook)
    receipt_path = Path(args.receipt)

    packet = load_json(packet_path)
    packet_sha256 = verify_packet(packet, packet_path, workbook_path, receipt_path)
    writes, by_fragment = build_write_maps(packet)
    terminal_sink_handle, non_self_source_handle = validate_single_terminal_non_self_binding(packet, by_fragment)
    validate_other_checked_sinks_self_bound(packet, by_fragment, writes)

    if terminal_sink_handle not in {item["sink_handle"] for item in writes}:
        raise ValueError("Terminal sink write is missing from resolved_cell_writes")

    original_sheetnames, original_widths = apply_writes(workbook_path, writes)
    recalc_meta = recalc_workbook(workbook_path)
    error_cells = scan_formula_errors(workbook_path)
    zero_formula_errors = not error_cells
    format_checks = verify_format_preservation(workbook_path, original_sheetnames, original_widths)

    expected_checks = {
        "no_extra_sheets",
        "task_sheet_column_widths_preserved",
        "data_sheet_column_widths_preserved",
        "xlsx_extension_preserved",
    }
    if set(format_checks) != expected_checks:
        missing = ", ".join(sorted(expected_checks - set(format_checks)))
        raise RuntimeError(f"Format preservation checks failed: {missing}")
    if not zero_formula_errors:
        joined = ", ".join(error_cells)
        raise RuntimeError(f"Excel errors found: {joined}")

    written_sink_value = load_written_sink_value(workbook_path, terminal_sink_handle)
    receipt = {
        "current_stage": CURRENT_STAGE,
        "next_required_skill": None,
        "allowed_outputs": [
            str(workbook_path.as_posix()),
            str(receipt_path.as_posix()),
        ],
        "packet_artifact": str(packet_path.as_posix()),
        "packet_path": str(packet_path.as_posix()),
        "packet_sha256": packet_sha256,
        "packet_digest": packet_sha256,
        "terminal_sink_handle": terminal_sink_handle,
        "non_self_source_handle": non_self_source_handle,
        "written_sink_value": written_sink_value,
        "workbook_path": str(workbook_path.as_posix()),
        "recalc_status": recalc_meta.get("status", "unknown"),
        "zero_formula_errors": zero_formula_errors,
        "format_preservation_checks": [
            "no_extra_sheets",
            "task_sheet_column_widths_preserved",
            "data_sheet_column_widths_preserved",
            "xlsx_extension_preserved",
        ],
    }
    write_receipt(receipt_path, receipt)


if __name__ == "__main__":
    main()
