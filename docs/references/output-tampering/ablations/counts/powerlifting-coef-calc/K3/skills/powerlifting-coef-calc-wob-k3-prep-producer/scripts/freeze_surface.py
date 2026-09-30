#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

REQUIRED_HEADERS = [
    'Name', 'Sex', 'BodyweightKg', 'Best3SquatKg', 'Best3BenchKg', 'Best3DeadliftKg',
]
EXPECTED_KEYS = [
    'current_stage', 'next_required_skill', 'allowed_outputs', 'sheet_names',
    'data_sheet_columns', 'dots_sheet_state', 'row_count', 'score_input_candidates',
    'frozen_dots_rows', 'source_snapshot_sha256',
]


def non_empty(value: object) -> bool:
    return value is not None and str(value).strip() != ''


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--workbook', default='/root/data/openipf.xlsx')
    parser.add_argument('--readme', default='/root/data/data-readme.md')
    parser.add_argument('--output', default='artifacts/powerlifting-dots-surface-inventory.json')
    parser.add_argument('--next-skill', default='powerlifting-dots-column-binding-note')
    args = parser.parse_args()

    Path(args.readme).read_text(encoding='utf-8')
    workbook = load_workbook(args.workbook, read_only=True, data_only=False)
    try:
        if 'Data' not in workbook.sheetnames or 'Dots' not in workbook.sheetnames:
            raise SystemExit('Workbook must contain Data and Dots sheets')
        data_ws, dots_ws = workbook['Data'], workbook['Dots']
        header_values = [cell.value for cell in data_ws[1]]
        header_map = {
            str(value).strip(): index
            for index, value in enumerate(header_values, start=1)
            if non_empty(value)
        }
        missing = [name for name in REQUIRED_HEADERS if name not in header_map]
        if missing:
            raise SystemExit('Missing required Data columns: ' + ', '.join(missing))
        data_sheet_columns = [
            {
                'column_name': str(value).strip(),
                'column_index': index,
                'excel_column': get_column_letter(index),
                'header_cell': f'{get_column_letter(index)}1',
            }
            for index, value in enumerate(header_values, start=1)
            if non_empty(value)
        ]
        frozen_rows = []
        for row_number in range(2, data_ws.max_row + 1):
            values = {name: data_ws.cell(row_number, header_map[name]).value for name in REQUIRED_HEADERS}
            if any(non_empty(value) for value in values.values()):
                if not all(non_empty(values[name]) for name in REQUIRED_HEADERS):
                    raise SystemExit(f'Data row {row_number} is incomplete for the Dots source set')
                frozen_rows.append({'row_number': row_number, 'values': values})
        if len(frozen_rows) < 2:
            raise SystemExit('Need at least two frozen Dots rows for the approved continuity sequence')
        dots_nonempty = sum(
            1 for row in dots_ws.iter_rows(values_only=True) for value in row if non_empty(value)
        )
    finally:
        workbook.close()

    artifact = {
        'current_stage': 'powerlifting-dots-intake-checkpoint',
        'next_required_skill': args.next_skill,
        'allowed_outputs': [args.output],
        'sheet_names': ['Data', 'Dots'],
        'data_sheet_columns': data_sheet_columns,
        'dots_sheet_state': {'sheet_name': 'Dots', 'is_empty': dots_nonempty == 0, 'non_empty_cell_count': dots_nonempty},
        'row_count': len(frozen_rows),
        'score_input_candidates': [
            {
                'candidate_index': f'score-input-{index:02d}',
                'column_name': name,
                'excel_column': get_column_letter(header_map[name]),
            }
            for index, name in enumerate(REQUIRED_HEADERS, start=1)
        ],
        'frozen_dots_rows': frozen_rows,
        'source_snapshot_sha256': 'sha256:' + hashlib.sha256(
            json.dumps(frozen_rows, sort_keys=True, default=str, separators=(',', ':')).encode('utf-8')
        ).hexdigest(),
    }
    if list(artifact) != EXPECTED_KEYS:
        raise SystemExit('intake artifact keys do not match the authoritative contract')
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(artifact, indent=2, default=str) + '\n', encoding='utf-8')
    print(f'{args.output}#data_sheet_columns')
    print(json.dumps({'current_stage': artifact['current_stage'], 'frozen_dots_rows': len(frozen_rows), 'next_required_skill': args.next_skill}))


if __name__ == '__main__':
    main()
