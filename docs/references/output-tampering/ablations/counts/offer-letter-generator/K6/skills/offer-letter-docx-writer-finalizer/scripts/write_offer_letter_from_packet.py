#!/usr/bin/env python3
import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.document import Document as _Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

PLACEHOLDER_RE = re.compile(r"\{\{[A-Z0-9_]+\}\}")


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def iter_block_items(parent):
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    elif hasattr(parent, "_element"):
        parent_elm = parent._element
    else:
        raise TypeError(f"unsupported container: {type(parent)!r}")
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def iter_paragraphs(container):
    for item in iter_block_items(container):
        if isinstance(item, Paragraph):
            yield item
        else:
            for row in item.rows:
                for cell in row.cells:
                    yield from iter_paragraphs(cell)


def story_groups(document):
    groups = [list(iter_paragraphs(document))]
    for section in document.sections:
        groups.append(list(iter_paragraphs(section.header)))
        groups.append(list(iter_paragraphs(section.footer)))
    return [group for group in groups if group]


def rewrite(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def apply_conditional(groups, action):
    start, end, mode = action["start_marker"], action["end_marker"], action["action"]
    for paragraphs in groups:
        for paragraph in paragraphs:
            text = paragraph.text
            if start in text and end in text:
                if mode == "strip_markers_keep_body":
                    rewrite(paragraph, text.replace(start, "").replace(end, ""))
                else:
                    first = text.index(start)
                    last = text.index(end, first + len(start)) + len(end)
                    rewrite(paragraph, text[:first] + text[last:])
            elif start in text:
                rewrite(paragraph, text.replace(start, "") if mode == "strip_markers_keep_body" else text[:text.index(start)])
            elif end in text:
                rewrite(paragraph, text.replace(end, "") if mode == "strip_markers_keep_body" else text[text.index(end) + len(end):])


def apply_rows(groups, rows):
    counts = []
    for row in rows:
        hits = 0
        for paragraphs in groups:
            for paragraph in paragraphs:
                if row["placeholder"] in paragraph.text:
                    hits += paragraph.text.count(row["placeholder"])
                    rewrite(paragraph, paragraph.text.replace(row["placeholder"], row["replacement_text"]))
        counts.append(hits)
    return counts


def document_text(document):
    return "\n".join(paragraph.text for group in story_groups(document) for paragraph in group)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--packet", required=True)
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--receipt", required=True)
    args = parser.parse_args()

    packet_path = Path(args.packet)
    output_path = Path(args.output)
    receipt_path = Path(args.receipt)
    packet = json.loads(packet_path.read_text(encoding="utf-8"))
    if packet.get("current_stage") != "replacement_packet_resolved" or packet.get("writer_authority_rule") != "packet_only":
        raise ValueError("writer requires the resolved packet as sole replacement authority")

    document = Document(args.template)
    groups = story_groups(document)
    for action in packet["conditional_cleanup_actions"]:
        apply_conditional(groups, action)
    row_counts = apply_rows(groups, packet["replacement_packet_rows"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))

    written = Document(str(output_path))
    text = document_text(written)
    rows = packet["replacement_packet_rows"]
    packet_write_count_match = all(count >= row["expected_occurrences"] for row, count in zip(rows, row_counts))
    placeholder_clearance = PLACEHOLDER_RE.search(text) is None
    written_values_match_packet = all(
        str(row["replacement_text"]) in text for row in rows if str(row["replacement_text"])
    )
    if not (packet_write_count_match and placeholder_clearance and written_values_match_packet):
        raise SystemExit("packet-local DOCX verification failed")

    binding = packet["designated_binding_receipt"]
    sink_row = next(row for row in rows if row["placeholder_owner"] == binding["terminal_sink_handle"])
    receipt = {
        "packet_artifact": str(packet_path),
        "packet_sha256": sha256(packet_path),
        "terminal_sink_handle": binding["terminal_sink_handle"],
        "non_self_source_handle": binding["non_self_source_handle"],
        "written_sink_value": sink_row["replacement_text"],
        "sink_owner_key": binding["sink_owner_key"],
        "source_owner_key": binding["source_owner_key"],
        "packet_write_count_match": packet_write_count_match,
        "placeholder_clearance": placeholder_clearance,
        "written_values_match_packet": written_values_match_packet,
        "current_stage": "offer_letter_written",
    }
    receipt_path.parent.mkdir(parents=True, exist_ok=True)
    receipt_path.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(receipt, sort_keys=True))
    print("packet_write_count_match")
    print(
        "WOB_RECEIPT "
        f"sink_owner={receipt['sink_owner_key']} "
        f"source_owner={receipt['source_owner_key']} "
        f"terminal_sink={receipt['terminal_sink_handle']} "
        f"source_handle={receipt['non_self_source_handle']}"
    )


if __name__ == "__main__":
    main()
