#!/usr/bin/env python3
"""Record the live workbook surface used by the weighted-GDP helpers."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


CURRENT_STAGE = "weighted-gdp-calc-workbook-intake-checkpoint"
NEXT_STAGE = "weighted-gdp-calc-workbook-summary-binding-note"
RANGES = {
    "lookup_formula_rows": "Task!H12:L31",
    "net_exports_formula_rows": "Task!H35:L40",
    "summary_stat_formula_rows": "Task!H42:L47",
    "weighted_mean_packet_rows": "Task!H50:L50",
}


def find_year_header_row(ws, fallback: int) -> int:
    for row in range(1, min(ws.max_row, 20) + 1):
        values = [ws.cell(row, col).value for col in range(8, 13)]
        numeric = [value for value in values if isinstance(value, (int, float))]
        if len(numeric) >= 3 and all(1900 <= value <= 2100 for value in numeric):
            return row
    return fallback


def main(argv: list[str]) -> None:
    if len(argv) != 3:
        raise SystemExit("usage: create_workbook_checkpoint.py <gdp.xlsx> <checkpoint.json>")

    workbook_path = Path(argv[1])
    output_path = Path(argv[2])
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Task" not in wb.sheetnames or "Data" not in wb.sheetnames:
            raise SystemExit("gdp.xlsx must contain Task and Data sheets")
        task_ws = wb["Task"]
        data_ws = wb["Data"]
        task_year_row = find_year_header_row(task_ws, fallback=9)
        data_year_row = find_year_header_row(data_ws, fallback=4)
        task_headers = [task_ws.cell(task_year_row, col).value for col in range(8, 13)]
        document = {
            "current_stage": CURRENT_STAGE,
            "next_required_skill": NEXT_STAGE,
            "allowed_outputs": ["artifacts/weighted_gdp_workbook_checkpoint.json"],
            "workbook_path": str(workbook_path),
            "sheet_names": list(wb.sheetnames),
            "observed_range_handles": RANGES,
            "observed_task_ranges": list(RANGES.values()),
            "task_year_header_row": task_year_row,
            "data_year_header_row": data_year_row,
            "observed_task_year_headers": task_headers,
            "source_data_rows": "Data!21:40",
            "designated_sink_target": "Task!J50",
            "formatting_guardrails": [
                "no_extra_sheets",
                "preserve_existing_column_widths",
                "preserve_existing_colors_fonts",
            ],
        }
    finally:
        wb.close()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(document, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main(sys.argv)
