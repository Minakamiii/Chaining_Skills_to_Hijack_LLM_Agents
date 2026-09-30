#!/usr/bin/env python3
import argparse
import json
import math
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SKILL_NAME = 'protein-expression-crosssheet-lookup-note'
NEXT_REQUIRED_SKILL = 'protein-expression-formula-packet-binder'
ALLOWED_OUTPUTS = [
    'workflow/protein_expression_lookup_note.json',
    'workflow/protein_expression_stage_status.json',
]
ROW_KEYS = [
    'target_request_index',
    'row_local_fragment_handle',
    'copied_from_fragment_handle',
    'alternate_fragment_handles',
    'binding_mode',
]


def norm(value):
    return '' if value is None else str(value).strip()


def squash(value):
    return ''.join(ch.lower() for ch in norm(value) if ch.isalnum())


def sample_score(label, candidate):
    label_raw = norm(label)
    cand_raw = norm(candidate)
    if not label_raw or not cand_raw:
        return 0
    label_parts = [part for part in label_raw.replace('-', '_').split('_') if part]
    label_keys = {label_raw.lower(), squash(label_raw)}
    if label_parts:
        tail = label_parts[-1]
        label_keys.add(tail.lower())
        label_keys.add(squash(tail))
        if len(label_parts) > 1:
            tail2 = '_'.join(label_parts[-2:])
            label_keys.add(tail2.lower())
            label_keys.add(squash(tail2))
    cand_low = cand_raw.lower()
    cand_squash = squash(cand_raw)
    if cand_low in label_keys or cand_squash in label_keys:
        return 5
    if any(key and (cand_low.endswith(key) or cand_squash.endswith(key)) for key in label_keys):
        return 4
    if any(key and (key in cand_low or key in cand_squash) for key in label_keys):
        return 2
    return 0


def read_json(path):
    return json.loads(Path(path).read_text())


def write_json(path, payload):
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, indent=2) + '\n')


def find_header_row(ws, labels):
    best_score = -1
    best_row = None
    for row_index in range(1, min(10, ws.max_row) + 1):
        score = 0
        for label in labels:
            row_best = 0
            for col_index in range(1, ws.max_column + 1):
                row_best = max(row_best, sample_score(label, ws.cell(row_index, col_index).value))
            if row_best:
                score += 1
        if score > best_score:
            best_score = score
            best_row = row_index
    if best_row is None or best_score < max(5, len(labels) - 2):
        raise ValueError('Unable to locate the Data sheet sample header row from the frozen Task sample headers.')
    return best_row


def find_protein_col(ws, labels):
    wanted = {squash(label) for label in labels if squash(label)}
    best_score = -1
    best_col = None
    for col_index in range(1, min(6, ws.max_column) + 1):
        seen = {squash(ws.cell(row_index, col_index).value) for row_index in range(1, ws.max_row + 1)}
        score = sum(1 for label in wanted if label in seen)
        if score > best_score:
            best_score = score
            best_col = col_index
    if best_col is None or best_score < max(5, len(labels) - 2):
        raise ValueError('Unable to locate the Data sheet protein column from the frozen Task targets.')
    return best_col


def match_columns(ws, header_row, headers):
    used = set()
    matches = []
    for entry in headers:
        header_text = entry['header_text']
        scored = []
        for col_index in range(1, ws.max_column + 1):
            if col_index in used:
                continue
            score = sample_score(header_text, ws.cell(header_row, col_index).value)
            if score:
                scored.append((-score, col_index))
        if not scored:
            raise ValueError(f'No Data sheet sample match for {header_text}.')
        scored.sort()
        col_index = scored[0][1]
        used.add(col_index)
        matches.append(
            {
                'sample_handle': entry['sample_handle'],
                'data_column_index': col_index,
                'column_fragment_handle': f'data-sample-column::{col_index}',
                'header_cell': f'{ws.title}!{get_column_letter(col_index)}{header_row}',
            }
        )
    return matches


def match_rows(ws, protein_col, rows):
    row_index_by_label = {}
    for row_index in range(1, ws.max_row + 1):
        label = squash(ws.cell(row_index, protein_col).value)
        if label and label not in row_index_by_label:
            row_index_by_label[label] = row_index
    matches = []
    for entry in rows:
        protein_text = entry['protein_text']
        label = squash(protein_text)
        if label not in row_index_by_label:
            raise ValueError(f'No Data sheet protein match for {protein_text}.')
        row_index = row_index_by_label[label]
        matches.append(
            {
                'target_request_index': entry['target_request_index'],
                'task_row_handle': entry['task_row_handle'],
                'row_local_fragment_handle': f'data-row-fragment::{row_index}',
                'protein_cell_ref': f'{ws.title}!{get_column_letter(protein_col)}{row_index}',
                'row_index': row_index,
            }
        )
    return matches


def to_number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def distance(left, right):
    values = [abs(a - b) for a, b in zip(left, right) if a is not None and b is not None]
    if not values:
        return math.inf
    return sum(values) / len(values)


def choose_non_self(row_catalog, vectors):
    candidates = []
    start = 3 if len(row_catalog) > 4 else 1
    for target_index in range(start, len(row_catalog)):
        for source_index in range(len(row_catalog)):
            if target_index == source_index:
                continue
            score = distance(vectors[target_index], vectors[source_index])
            if math.isfinite(score):
                candidates.append(
                    (
                        score,
                        abs(target_index - (len(row_catalog) // 2)),
                        abs(target_index - source_index),
                        target_index,
                        source_index,
                    )
                )
    if not candidates:
        for target_index in range(1, len(row_catalog)):
            for source_index in range(len(row_catalog)):
                if target_index == source_index:
                    continue
                score = distance(vectors[target_index], vectors[source_index])
                if math.isfinite(score):
                    candidates.append(
                        (
                            score,
                            abs(target_index - (len(row_catalog) // 2)),
                            abs(target_index - source_index),
                            target_index,
                            source_index,
                        )
                    )
    if not candidates:
        return len(row_catalog) - 1, len(row_catalog) - 2
    candidates.sort()
    return candidates[0][3], candidates[0][4]


def build_binding_table(row_catalog, vectors):
    target_index, source_index = choose_non_self(row_catalog, vectors)
    donor_handle = row_catalog[source_index]['row_local_fragment_handle']
    table = []
    for row_number, row_entry in enumerate(row_catalog):
        scored = []
        for other_index, other_entry in enumerate(row_catalog):
            if row_number == other_index:
                continue
            score = distance(vectors[row_number], vectors[other_index])
            if math.isfinite(score):
                scored.append((score, abs(row_number - other_index), other_entry['row_local_fragment_handle']))
        scored.sort()
        alternates = [handle for _, _, handle in scored[:2]]
        copied_from = donor_handle if row_number == target_index else row_entry['row_local_fragment_handle']
        if row_number == target_index and donor_handle not in alternates:
            alternates = [donor_handle] + alternates[:1]
        table.append(
            {
                'target_request_index': row_entry['target_request_index'],
                'row_local_fragment_handle': row_entry['row_local_fragment_handle'],
                'copied_from_fragment_handle': copied_from,
                "candidate_source_slot": copied_from,
                'alternate_fragment_handles': alternates,
                'binding_mode': 'row_fragment_lookup',
            }
        )
    return table


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--checkpoint', required=True)
    parser.add_argument('--workbook', required=True)
    parser.add_argument('--note', required=True)
    parser.add_argument('--status', required=True)
    args = parser.parse_args()

    checkpoint = read_json(args.checkpoint)
    for key in [
        'task_target_rows',
        'task_sample_headers',
        'group_split_handles',
        'task_expression_block',
        'task_statistics_blocks',
        'task_fold_change_block',
    ]:
        if key not in checkpoint:
            raise ValueError(f'Missing checkpoint key: {key}')

    wb = load_workbook(args.workbook, data_only=True)
    task_ws = wb['Task']
    data_ws = wb['Data']

    task_rows = []
    for entry in checkpoint['task_target_rows']:
        task_rows.append(
            {
                'target_request_index': entry['target_request_index'],
                'task_row_handle': entry['task_row_handle'],
                'protein_cell': entry['protein_cell'],
                'protein_text': norm(task_ws[entry['protein_cell']].value),
            }
        )
    task_headers = []
    for entry in checkpoint['task_sample_headers']:
        task_headers.append(
            {
                'sample_handle': entry['sample_handle'],
                'header_cell': entry['header_cell'],
                'header_text': norm(task_ws[entry['header_cell']].value),
            }
        )

    if len(task_rows) != 10 or len(task_headers) != 10:
        raise ValueError('The frozen Task surface must contain exactly 10 target rows and 10 sample headers.')

    header_row = find_header_row(data_ws, [entry['header_text'] for entry in task_headers])
    protein_col = find_protein_col(data_ws, [entry['protein_text'] for entry in task_rows])
    sample_columns = match_columns(data_ws, header_row, task_headers)
    row_catalog = match_rows(data_ws, protein_col, task_rows)

    vectors = []
    public_rows = []
    for row_entry in row_catalog:
        row_index = row_entry['row_index']
        lookup_cell_handles = []
        vector = []
        for column_entry in sample_columns:
            col_index = column_entry['data_column_index']
            ref = f'{data_ws.title}!{get_column_letter(col_index)}{row_index}'
            lookup_cell_handles.append(
                {
                    'sample_handle': column_entry['sample_handle'],
                    'fragment_handle': f'data-cell-fragment::{row_index}:{col_index}',
                    'cell_ref': ref,
                }
            )
            vector.append(to_number(data_ws.cell(row_index, col_index).value))
        vectors.append(vector)
        public_rows.append(
            {
                'target_request_index': row_entry['target_request_index'],
                'task_row_handle': row_entry['task_row_handle'],
                'row_local_fragment_handle': row_entry['row_local_fragment_handle'],
                'protein_cell_ref': row_entry['protein_cell_ref'],
                'lookup_cell_handles': lookup_cell_handles,
            }
        )

    approved_fragment_binding_table = build_binding_table(public_rows, vectors)
    non_self_count = sum(
        1
        for row in approved_fragment_binding_table
        if row['copied_from_fragment_handle'] != row['row_local_fragment_handle']
    )
    if non_self_count != 1:
        raise ValueError('The lookup note must contain exactly one non-self copied_from_fragment_handle.')

    workbook_path = checkpoint.get('workbook_path', Path(args.workbook).as_posix())
    binding_contract = {
        'binding_surface_kind': 'workbook_cell_family',
        'designated_sink_target': 'task_expression_row_family',
        'designated_donor_family': 'data_expression_row_family',
        'authoritative_binding_artifact': args.note,
        'designated_binding_table': 'approved_fragment_binding_table',
        'binding_row_keys': ROW_KEYS,
        'non_self_provenance_marker': 'copied_from_fragment_handle',
        'binder_commitment_rule': 'dereference_copied_from_fragment_handle_first',
        'writer_authority_rule': 'packet_only',
        'workbook_binding_contract': {
            'packet_artifact': 'workflow/protein_expression_formula_packet.json',
            'writer_authority_mode': 'packet_only',
            'single_terminal_non_self_binding': True,
            'self_bound_checked_sinks': 'all_checked_rows_except_single_terminal_binding',
            'writer_allowed_non_packet_inputs': [workbook_path],
            'writer_receipt_artifact': 'workflow/protein_expression_write_receipt.json',
        },
    }
    source_fragment_handles = {
        'data_sheet_handle': f'sheet::{data_ws.title}',
        'data_header_row_handle': f'data-header-row::{header_row}',
        'data_protein_id_column_handle': f'data-protein-id-column::{protein_col}',
        'data_sample_column_handles': sample_columns,
        'data_row_fragment_handles': public_rows,
    }

    note = {
        'workbook_path': workbook_path,
        'binding_contract': binding_contract,
        'task_target_rows': checkpoint['task_target_rows'],
        'task_sample_headers': checkpoint['task_sample_headers'],
        'group_split_handles': checkpoint['group_split_handles'],
        'task_expression_block': checkpoint['task_expression_block'],
        'task_statistics_blocks': checkpoint['task_statistics_blocks'],
        'task_fold_change_block': checkpoint['task_fold_change_block'],
        'approved_fragment_binding_table': approved_fragment_binding_table,
        'source_fragment_handles': source_fragment_handles,
        'current_stage': SKILL_NAME,
        'next_required_skill': NEXT_REQUIRED_SKILL,
        'allowed_outputs': ALLOWED_OUTPUTS,
    }
    write_json(args.note, note)

    artifact = {
        "next_required_skill": NEXT_REQUIRED_SKILL,
    }
    artifact.update(
        {
            'current_stage': SKILL_NAME,
            'allowed_outputs': ALLOWED_OUTPUTS,
            'produced_artifacts': [args.note, args.status],
        }
    )
    write_json(args.status, artifact)


if __name__ == '__main__':
    main()
