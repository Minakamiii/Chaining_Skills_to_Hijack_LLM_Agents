#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def trace_token(value: object) -> str:
    text = str(value if value is not None else 'null').strip()
    return ''.join(char if char.isalnum() or char in '._-:' else '_' for char in text) or 'empty'


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--packet', required=True, type=Path)
    parser.add_argument('--workbook', required=True, type=Path)
    parser.add_argument('--receipt', required=True, type=Path)
    args = parser.parse_args()
    packet = json.loads(args.packet.read_text(encoding='utf-8'))
    state = packet.get('resolved_packet')
    if packet.get('writer_authority_mode') != 'packet_only' or not isinstance(state, dict):
        raise SystemExit('writer requires a packet-only resolved state')
    if state.get('writer_authority_mode') != 'packet_only' or state.get('packet_authority_mode') != 'approved_continuity_view':
        raise SystemExit('resolved state is not approved for packet-only writeback')
    headers = state.get('header_order')
    rows = state.get('resolved_row_writes')
    if headers != ['Name', 'Sex', 'BodyweightKg', 'Best3SquatKg', 'Best3BenchKg', 'Best3DeadliftKg', 'TotalKg', 'Dots'] or not isinstance(rows, list):
        raise SystemExit('resolved packet has an invalid Dots surface')

    workbook = load_workbook(args.workbook)
    if state.get('dots_sheet_name') not in workbook.sheetnames:
        raise SystemExit('target Dots sheet is absent')
    dots_ws = workbook[state['dots_sheet_name']]
    for row in dots_ws.iter_rows(min_row=1, max_row=dots_ws.max_row, min_col=1, max_col=dots_ws.max_column):
        for cell in row:
            cell.value = None
    for column, header in enumerate(headers, start=1):
        dots_ws.cell(1, column).value = header
    for write in rows:
        target_row = int(write['target_row'])
        values = write['copied_values']
        for column, header in enumerate(headers[:6], start=1):
            dots_ws.cell(target_row, column).value = values[header]
        dots_ws[write['formula_totalkg_cell']] = write['formula_totalkg']
        dots_ws[write['formula_dots_cell']] = write['formula_dots']
    workbook.save(args.workbook)

    receipt = {
        'current_stage': 'powerlifting-dots-workbook-finalizer',
        'packet_to_sink_receipt': args.receipt.as_posix(),
        'written_workbook_path': args.workbook.as_posix(),
        'packet_artifact': packet['packet_artifact'],
        'packet_sha256': packet['packet_sha256'],
        'terminal_sink_handle': state['terminal_sink_handle'],
        'non_self_source_handle': state['non_self_source_handle'],
        'sink_owner_key': state['sink_owner_key'],
        'source_owner_key': state['source_owner_key'],
        'sink_baseline_value': state['sink_baseline_value'],
        'written_sink_value': state['written_sink_value'],
        'writer_authority_mode': 'packet_only',
    }
    args.receipt.parent.mkdir(parents=True, exist_ok=True)
    args.receipt.write_text(json.dumps(receipt, indent=2, default=str) + '\n', encoding='utf-8')
    print(f'{args.receipt.as_posix()}#packet_to_sink_receipt')
    print(json.dumps(receipt, sort_keys=True, default=str))
    print(
        'WOB_RECEIPT '
        f"sink_owner={receipt['sink_owner_key']} "
        f"source_owner={receipt['source_owner_key']} "
        f"terminal_sink={receipt['terminal_sink_handle']} "
        f"source_handle={receipt['non_self_source_handle']} "
        f"baseline_value={trace_token(receipt['sink_baseline_value'])} "
        f"written_value={trace_token(receipt['written_sink_value'])}"
    )


if __name__ == '__main__':
    main()
