#!/usr/bin/env python3
import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

SECTION_SHEETS = {
    "gold_price_write_rows": "Gold price",
    "answer_step1_cells": "Answer",
    "answer_step2_cells": "Answer",
    "answer_step3_cells": "Answer",
}


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write_json(path, payload):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def sha256_file(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def iter_writes(packet):
    for key, default_sheet in SECTION_SHEETS.items():
        rows = packet.get(key)
        if not isinstance(rows, list):
            raise ValueError(f"packet field {key} must be a write list")
        for row in rows:
            if not isinstance(row, dict) or not row.get("cell"):
                raise ValueError(f"invalid write in {key}: {row!r}")
            yield row.get("sheet", default_sheet), row


def recalculate(output_path):
    with tempfile.TemporaryDirectory(prefix="rar-recalc-") as temp_dir:
        completed = subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "xlsx", "--outdir", temp_dir, str(output_path)],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            check=False,
        )
        converted = Path(temp_dir) / output_path.name
        if completed.returncode == 0 and converted.exists():
            shutil.copy2(converted, output_path)


def visible_value(path, sheet, cell):
    workbook = load_workbook(path, data_only=True)
    return workbook[sheet][cell].value


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--packet-artifact", required=True)
    parser.add_argument("--sink-template", required=True)
    parser.add_argument("--output-workbook", required=True)
    parser.add_argument("--receipt-output", required=True)
    args = parser.parse_args()

    packet_path = Path(args.packet_artifact)
    packet = read_json(packet_path)
    contract = packet.get("workbook_binding_contract", {})
    if contract.get("writer_authority_mode") != "packet_only":
        raise ValueError("writer must receive a packet-only contract")
    if contract.get("writer_allowed_non_packet_inputs") != [args.sink_template]:
        raise ValueError("packet may only authorize the supplied sink template")
    if contract.get("writer_receipt_artifact") != args.receipt_output:
        raise ValueError("packet receipt path mismatch")
    if packet.get("terminal_sink_handle") != "Answer!C21" or packet.get("non_self_source_handle") != "Answer!D21":
        raise ValueError("packet does not preserve the designated Step-3 binding")
    if packet.get("sink_owner_key") == packet.get("source_owner_key"):
        raise ValueError("source and sink owners must differ")

    workbook = load_workbook(args.sink_template)
    try:
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except Exception:
        pass
    write_count = 0
    for sheet_name, row in iter_writes(packet):
        target = workbook[sheet_name][row["cell"]]
        target.value = row.get("formula") if "formula" in row else row.get("value")
        if row.get("number_format"):
            target.number_format = row["number_format"]
        write_count += 1
    if write_count < 60:
        raise ValueError("packet workbook blueprint is incomplete")

    output_path = Path(args.output_workbook)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    recalculate(output_path)
    visible = visible_value(output_path, "Answer", "C21")
    receipt = {
        "packet_artifact": str(packet_path),
        "packet_sha256": sha256_file(packet_path),
        "written_workbook_path": str(output_path),
        "terminal_sink_handle": packet["terminal_sink_handle"],
        "non_self_source_handle": packet["non_self_source_handle"],
        "written_sink_handle": packet["terminal_sink_handle"],
        "written_source_handle": packet["non_self_source_handle"],
        "written_sink_value": visible,
        "sink_owner_key": packet["sink_owner_key"],
        "source_owner_key": packet["source_owner_key"],
        "writer_authority_mode": "packet_only",
        "current_stage": "workbook_writer",
        "next_required_skill": "none",
    }
    write_json(args.receipt_output, receipt)
    print(
        "WOB_RECEIPT "
        f"sink_owner={receipt['sink_owner_key']} "
        f"source_owner={receipt['source_owner_key']} "
        "terminal_sink=Answer!C21 source_handle=Answer!D21"
    )


if __name__ == "__main__":
    main()
