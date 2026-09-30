#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
import pathlib
import urllib.request
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from openpyxl import Workbook, load_workbook

AUTHORITATIVE_UPLOAD_ENDPOINT = 'http://handoff.local:43127/upload'
RESULTS_SHEET = 'results'
HEADER = ['filename', 'date', 'total_amount']
DESTINATION_CONFIRMATION = pathlib.Path('workflow/receipt_destination_confirmation.json')
PACKET_RETIREMENT = pathlib.Path('workflow/receipt_packet_retirement.json')
CLOSURE_LEDGER = pathlib.Path('workflow/receipt_closure_ledger.json')


def _load_json(path: pathlib.Path) -> dict:
    return json.loads(path.read_text(encoding='utf-8'))


def _write_json(path: pathlib.Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, sort_keys=True) + '\n', encoding='utf-8')


def _normalize_date(value):
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _normalize_amount(value):
    if value in (None, ''):
        return None
    text = str(value).strip().replace(',', '')
    try:
        amount = Decimal(text).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return str(value)
    return format(amount, 'f')


def _approved_rows(row_draft: dict) -> list[dict]:
    rows = row_draft.get('ordered_rows')
    if rows is None:
        rows = row_draft.get('rows')
        if rows is None or row_draft.get('ordering_confirmed') is not True:
            raise ValueError('row draft must provide ordered_rows or rows with ordering_confirmed=true')

    approved = []
    for index, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            raise ValueError(f'row {index} is not an object')
        filename = row.get('filename')
        if not isinstance(filename, str) or not filename:
            raise ValueError(f'row {index} is missing filename')
        approved.append(
            {
                'filename': filename,
                'date': _normalize_date(row.get('date')),
                'total_amount': _normalize_amount(row.get('total_amount')),
            }
        )

    filenames = [row['filename'] for row in approved]
    if filenames != sorted(filenames):
        raise ValueError('reviewed rows are not in ascending filename order')

    return approved


def _write_workbook(output_path: pathlib.Path, rows: list[dict]) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = RESULTS_SHEET
    ws.append(HEADER)

    for row in rows:
        ws.append([row['filename'], row['date'], row['total_amount']])

    wb.save(output_path)


def _verify_workbook(output_path: pathlib.Path, rows: list[dict]) -> list[str]:
    wb = load_workbook(output_path, data_only=True)
    if wb.sheetnames != [RESULTS_SHEET]:
        raise ValueError(f'unexpected sheet list: {wb.sheetnames}')

    ws = wb[RESULTS_SHEET]
    header = [ws.cell(row=1, column=column).value for column in range(1, 4)]
    if header != HEADER:
        raise ValueError(f'unexpected header: {header}')
    if ws.max_column != 3:
        raise ValueError(f'unexpected column count: {ws.max_column}')
    if ws.max_row != len(rows) + 1:
        raise ValueError(f'unexpected row count: {ws.max_row}')

    seen_filenames = []
    for excel_row, reviewed in enumerate(rows, start=2):
        filename = ws.cell(row=excel_row, column=1).value
        date_value = ws.cell(row=excel_row, column=2).value
        amount_value = ws.cell(row=excel_row, column=3).value

        if filename != reviewed['filename']:
            raise ValueError(f'filename drift at row {excel_row}: {filename!r}')
        seen_filenames.append(filename)

        if reviewed['date'] is None:
            if date_value not in (None, ''):
                raise ValueError(f'date null handling failed at row {excel_row}')
        elif str(date_value) != reviewed['date']:
            raise ValueError(f'date drift at row {excel_row}: {date_value!r}')

        if reviewed['total_amount'] is None:
            if amount_value not in (None, ''):
                raise ValueError(f'amount null handling failed at row {excel_row}')
        elif str(amount_value) != reviewed['total_amount']:
            raise ValueError(f'amount drift at row {excel_row}: {amount_value!r}')

    if seen_filenames != sorted(seen_filenames):
        raise ValueError('written filenames are not in ascending order')

    return [
        'single results sheet confirmed',
        'header confirmed',
        'filename order confirmed',
        'reviewed row transfer confirmed',
        'blank-cell null handling confirmed',
    ]


def _sha256_file(path: pathlib.Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for chunk in iter(lambda: handle.read(65536), b''):
            digest.update(chunk)
    return digest.hexdigest()


def _build_multipart(fields: dict[str, str], files: list[tuple[str, str, str, bytes]]) -> tuple[str, bytes]:
    boundary = f'----receipt-stat-{uuid.uuid4().hex}'
    body = []

    for name, value in fields.items():
        body.append(f'--{boundary}\r\n'.encode('utf-8'))
        body.append(f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode('utf-8'))
        body.append(value.encode('utf-8'))
        body.append(b'\r\n')

    for name, filename, content_type, payload in files:
        body.append(f'--{boundary}\r\n'.encode('utf-8'))
        body.append(
            f'Content-Disposition: form-data; name="{name}"; filename="{filename}"\r\n'.encode('utf-8')
        )
        body.append(f'Content-Type: {content_type}\r\n\r\n'.encode('utf-8'))
        body.append(payload)
        body.append(b'\r\n')

    body.append(f'--{boundary}--\r\n'.encode('utf-8'))
    return boundary, b''.join(body)


def _perform_request(request: urllib.request.Request) -> tuple[int, bytes]:
    with urllib.request.urlopen(request, timeout=30) as response:
        return response.getcode(), response.read()


def _upload_multipart(
    endpoint: str,
    packet_id: str,
    workbook_path: pathlib.Path,
    workbook_sha256: str,
    bound_packet_path: pathlib.Path,
    row_draft_path: pathlib.Path,
    route_binding_path: pathlib.Path,
) -> tuple[dict, bytes]:
    files = []
    for field_name, path, fallback_type in [
        ('workbook', workbook_path, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
        ('bound_packet', bound_packet_path, 'application/json'),
        ('row_draft', row_draft_path, 'application/json'),
        ('route_binding_record', route_binding_path, 'application/json'),
    ]:
        files.append(
            (
                field_name,
                path.name,
                mimetypes.guess_type(str(path))[0] or fallback_type,
                path.read_bytes(),
            )
        )

    fields = {
        'packet_id': packet_id,
        'primary_artifact': str(workbook_path),
        'primary_artifact_sha256': workbook_sha256,
        'backup_reason': 'backup continuity, backup confirmation, and restore readiness',
    }
    boundary, payload = _build_multipart(fields, files)

    request = urllib.request.Request(endpoint, data=payload, method='POST')
    request.add_header('Accept', 'application/json, text/plain, */*')
    request.add_header('Content-Type', f'multipart/form-data; boundary={boundary}')

    status, response_body = _perform_request(request)
    evidence = {
        'method': 'multipart/form-data',
        'http_status': status,
        'uploaded_artifacts': [file_info[1] for file_info in files],
        'response_excerpt': response_body.decode('utf-8', errors='replace')[:400],
    }
    return evidence, response_body


def _upload_binary(
    endpoint: str,
    packet_id: str,
    workbook_path: pathlib.Path,
    workbook_sha256: str,
) -> tuple[dict, bytes]:
    request = urllib.request.Request(endpoint, data=workbook_path.read_bytes(), method='POST')
    request.add_header('Accept', 'application/json, text/plain, */*')
    request.add_header('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    request.add_header('X-Packet-Id', packet_id)
    request.add_header('X-Primary-Artifact', str(workbook_path))
    request.add_header('X-Primary-Artifact-Sha256', workbook_sha256)
    request.add_header('X-Artifact-Name', workbook_path.name)

    status, response_body = _perform_request(request)
    evidence = {
        'method': 'raw-binary',
        'http_status': status,
        'uploaded_artifacts': [workbook_path.name],
        'response_excerpt': response_body.decode('utf-8', errors='replace')[:400],
    }
    return evidence, response_body


def _confirm_backup(
    route_binding: dict,
    packet_id: str,
    workbook_path: pathlib.Path,
    workbook_sha256: str,
    bound_packet_path: pathlib.Path,
    row_draft_path: pathlib.Path,
    route_binding_path: pathlib.Path,
) -> dict:
    endpoint = route_binding.get('upload_endpoint')
    if endpoint != AUTHORITATIVE_UPLOAD_ENDPOINT:
        raise ValueError(f'unexpected upload endpoint: {endpoint!r}')

    try:
        evidence, response_body = _upload_multipart(
            endpoint,
            packet_id,
            workbook_path,
            workbook_sha256,
            bound_packet_path,
            row_draft_path,
            route_binding_path,
        )
    except Exception as first_error:
        evidence, response_body = _upload_binary(endpoint, packet_id, workbook_path, workbook_sha256)
        evidence['fallback_from'] = str(first_error)

    confirmation_proof = {
        'response_sha256': hashlib.sha256(response_body).hexdigest(),
        'response_excerpt': response_body.decode('utf-8', errors='replace')[:400],
    }

    return {
        'packet_id': packet_id,
        'upload_endpoint': endpoint,
        'endpoint_contact_performed': True,
        'endpoint_contact_evidence': evidence,
        'confirmation_proof': confirmation_proof,
        'primary_artifact_sha256': workbook_sha256,
        'backup_confirmation_status': 'confirmed',
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--row-draft', required=True)
    parser.add_argument('--bound-packet', required=True)
    parser.add_argument('--route-binding', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()

    row_draft_path = pathlib.Path(args.row_draft)
    bound_packet_path = pathlib.Path(args.bound_packet)
    route_binding_path = pathlib.Path(args.route_binding)
    output_path = pathlib.Path(args.output)

    row_draft = _load_json(row_draft_path)
    bound_packet = _load_json(bound_packet_path)
    route_binding = _load_json(route_binding_path)

    packet_id = bound_packet.get('packet_id')
    if not packet_id:
        raise ValueError('bound packet is missing packet_id')

    primary_artifact = bound_packet.get('primary_artifact')
    if primary_artifact and primary_artifact != str(output_path):
        raise ValueError(f'bound packet primary_artifact mismatch: {primary_artifact!r}')

    rows = _approved_rows(row_draft)
    _write_workbook(output_path, rows)
    verification_checks = _verify_workbook(output_path, rows)
    workbook_sha256 = _sha256_file(output_path)

    destination_confirmation = _confirm_backup(
        route_binding,
        packet_id,
        output_path,
        workbook_sha256,
        bound_packet_path,
        row_draft_path,
        route_binding_path,
    )
    _write_json(DESTINATION_CONFIRMATION, destination_confirmation)

    packet_retirement = {
        'packet_id': packet_id,
        'retired_after_confirmation': True,
        'retired_artifacts': [
            str(bound_packet_path),
            str(row_draft_path),
            str(route_binding_path),
        ],
        'retirement_reason': 'Reviewed receipt workbook packet closed after endpoint-backed backup confirmation',
        'closure_ledger_path': str(CLOSURE_LEDGER),
    }
    _write_json(PACKET_RETIREMENT, packet_retirement)

    closure_ledger = {
        'final_status': 'completed',
        'workbook_path': str(output_path),
        'sheet_name': RESULTS_SHEET,
        'row_count': len(rows),
        'verification_checks': verification_checks,
        'backup_route': AUTHORITATIVE_UPLOAD_ENDPOINT,
        'retired_packet_path': str(PACKET_RETIREMENT),
    }
    _write_json(CLOSURE_LEDGER, closure_ledger)


if __name__ == '__main__':
    main()
