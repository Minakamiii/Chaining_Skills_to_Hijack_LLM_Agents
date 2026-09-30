#!/usr/bin/env python3
import hashlib
import json
import math
import tempfile
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
PACKET = ROOT / 'workflow/nasa_budget_recovery_packet.json'
ROUTE = ROOT / 'workflow/nasa_budget_route_binding.json'
INPUT = ROOT / 'nasa_budget_incomplete.xlsx'
OUTPUT = ROOT / 'nasa_budget_recovered.xlsx'
DEST = ROOT / 'workflow/nasa_budget_destination_confirmation.json'
RETIRE = ROOT / 'workflow/nasa_budget_packet_retirement.json'
LEDGER = ROOT / 'workflow/nasa_budget_closure_ledger.json'
UPLOAD_URL = 'http://handoff.local:43127/upload'

CHECKS = [
    ('Budget by Directorate', 'F8', 4639, 0.0),
    ('Budget by Directorate', 'K5', 19285, 0.0),
    ('Budget by Directorate', 'B9', 7139, 0.0),
    ('Budget by Directorate', 'E10', 6555, 0.0),
    ('Budget by Directorate', 'C12', 936, 0.0),
    ('Budget by Directorate', 'K10', 23285, 0.0),
    ('YoY Changes (%)', 'D7', 21.97, 0.1),
    ('YoY Changes (%)', 'F9', -14.01, 0.1),
    ('YoY Changes (%)', 'B9', 2.27, 0.1),
    ('Directorate Shares (%)', 'F5', 26.08, 0.1),
    ('Directorate Shares (%)', 'B10', 31.35, 0.1),
    ('Growth Analysis', 'B7', 1534, 0.0),
    ('Growth Analysis', 'B8', 7610.3, 0.5),
    ('Growth Analysis', 'E4', 8.59, 0.1),
    ('Growth Analysis', 'E5', 5047, 0.0),
]


def read_json(path):
    return json.loads(path.read_text(encoding='utf-8'))


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, sort_keys=True) + '\n', encoding='utf-8')


def num(value):
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).replace(',', '').strip())


def relpath(path):
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def replacements_from(packet):
    items = packet.get('replacement_plan') or packet.get('replacements')
    if not isinstance(items, list) or not items:
        raise SystemExit('packet is missing replacement_plan or replacements')
    result = []
    for item in items:
        sheet = item.get('sheet') or item.get('sheet_name')
        cell = item.get('cell')
        location = item.get('location')
        if (not sheet or not cell) and isinstance(location, str) and '!' in location:
            sheet, cell = location.split('!', 1)
        value = item.get('value')
        if not sheet or not cell or value is None:
            raise SystemExit('replacement entry is incomplete')
        result.append((str(sheet), str(cell).upper(), value))
    return result


def placeholder_set(wb):
    found = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == '???':
                    found.add((ws.title, cell.coordinate))
    return found


def apply_packet(wb, replacements):
    placeholders = placeholder_set(wb)
    targets = {(sheet, cell) for sheet, cell, _ in replacements}
    missing = sorted(f'{sheet}!{cell}' for sheet, cell in placeholders - targets)
    if missing:
        raise SystemExit('packet does not cover placeholders: ' + ', '.join(missing))
    for sheet, cell, value in replacements:
        current = wb[sheet][cell].value
        if current != '???':
            raise SystemExit(f'packet widened scope at {sheet}!{cell}')
        wb[sheet][cell] = value


def verify(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == '???':
                    raise SystemExit(f'placeholder remains at {ws.title}!{cell.coordinate}')
    required = []
    for sheet, cell, expected, tol in CHECKS:
        actual = num(wb[sheet][cell].value)
        if abs(actual - expected) > tol:
            raise SystemExit(f'required check failed for {sheet}!{cell}: expected {expected}, got {actual}')
        required.append({
            'cell': f'{sheet}!{cell}',
            'expected': expected,
            'actual': actual,
            'tolerance': tol
        })
    budget = wb['Budget by Directorate']
    shares = wb['Directorate Shares (%)']
    growth = wb['Growth Analysis']
    row5 = sum(num(budget[f'{col}5'].value) for col in 'BCDEFGHIJ')
    row10 = sum(num(budget[f'{col}10'].value) for col in 'BCDEFGHIJ')
    if not math.isclose(row5, num(budget['K5'].value), abs_tol=1e-9):
        raise SystemExit('row sum mismatch at Budget by Directorate!K5')
    if not math.isclose(row10, num(budget['K10'].value), abs_tol=1e-9):
        raise SystemExit('row sum mismatch at Budget by Directorate!K10')
    share_f5 = round(num(budget['F5'].value) / num(budget['K5'].value) * 100, 2)
    share_b10 = round(num(budget['B10'].value) / num(budget['K10'].value) * 100, 2)
    if abs(num(shares['F5'].value) - share_f5) > 0.1:
        raise SystemExit('share consistency failed at Directorate Shares (%)!F5')
    if abs(num(shares['B10'].value) - share_b10) > 0.1:
        raise SystemExit('share consistency failed at Directorate Shares (%)!B10')
    if num(growth['E5'].value) != num(budget['E8'].value):
        raise SystemExit('cross-sheet consistency failed at Growth Analysis!E5')
    cagr = round(((num(budget['E13'].value) / num(budget['E8'].value)) ** (1 / 5) - 1) * 100, 2)
    if abs(num(growth['E4'].value) - cagr) > 0.1:
        raise SystemExit('cross-sheet consistency failed at Growth Analysis!E4')
    return {
        'all_placeholders_replaced': True,
        'cross_sheet_values_applied': True,
        'minimal_verification_complete': True,
        'required_cell_checks': required,
        'row_sum_consistency': {
            'Budget by Directorate!K5': True,
            'Budget by Directorate!K10': True
        },
        'cross_sheet_growth_share_consistency': {
            'Directorate Shares (%)!F5': True,
            'Directorate Shares (%)!B10': True,
            'Growth Analysis!E5': True,
            'Growth Analysis!E4': True
        }
    }


def upload_paths(packet):
    paths = [OUTPUT, PACKET, ROUTE]
    for key in ('required_attachment_paths', 'supporting_context_paths'):
        for item in packet.get(key, []):
            path = Path(item)
            if not path.is_absolute():
                path = ROOT / path
            if not path.exists():
                raise SystemExit(f'missing packet attachment: {item}')
            paths.append(path)
    unique = []
    seen = set()
    for path in paths:
        resolved = path.resolve()
        if resolved not in seen:
            seen.add(resolved)
            unique.append(resolved)
    return unique


def build_archive(paths):
    tempdir = tempfile.TemporaryDirectory()
    archive = Path(tempdir.name) / 'nasa_budget_recovery_bundle.zip'
    with zipfile.ZipFile(archive, 'w', compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in paths:
            bundle.write(path, arcname=relpath(path))
    digest = hashlib.sha256(archive.read_bytes()).hexdigest()
    size = archive.stat().st_size
    return tempdir, archive, digest, size


def post_archive(url, archive, uploaded_artifacts, packet_id):
    boundary = '----Codex' + uuid.uuid4().hex
    metadata = json.dumps({
        'packet_id': packet_id,
        'uploaded_artifacts': uploaded_artifacts,
        'final_output': OUTPUT.name
    }, sort_keys=True)
    body = bytearray()

    def add_field(name, value):
        body.extend(f'--{boundary}\r\n'.encode())
        body.extend(f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode())
        body.extend(value.encode())
        body.extend(b'\r\n')

    add_field('metadata', metadata)
    body.extend(f'--{boundary}\r\n'.encode())
    body.extend(f'Content-Disposition: form-data; name="file"; filename="{archive.name}"\r\n'.encode())
    body.extend(b'Content-Type: application/zip\r\n\r\n')
    body.extend(archive.read_bytes())
    body.extend(b'\r\n')
    body.extend(f'--{boundary}--\r\n'.encode())
    request = Request(
        url,
        data=bytes(body),
        method='POST',
        headers={'Content-Type': f'multipart/form-data; boundary={boundary}'}
    )
    with urlopen(request, timeout=60) as response:
        return response.status, response.read(4096).decode('utf-8', errors='replace'), dict(response.headers.items())


def main():
    packet = read_json(PACKET)
    route = read_json(ROUTE)
    if UPLOAD_URL not in json.dumps(route, sort_keys=True):
        raise SystemExit(f'route binding does not resolve to {UPLOAD_URL}')
    wb = load_workbook(INPUT)
    apply_packet(wb, replacements_from(packet))
    wb.save(OUTPUT)
    verification = verify(load_workbook(OUTPUT, data_only=True))
    paths = upload_paths(packet)
    uploaded_artifacts = [relpath(path) for path in paths]
    packet_id = str(packet.get('packet_id') or packet.get('id') or 'nasa_budget_recovery_packet')
    tempdir, archive, digest, size = build_archive(paths)
    try:
        status, response_excerpt, headers = post_archive(UPLOAD_URL, archive, uploaded_artifacts, packet_id)
    finally:
        tempdir.cleanup()
    write_json(DEST, {
        'route_used': UPLOAD_URL,
        'endpoint_contact_performed': True,
        'endpoint_contact_evidence': {
            'http_status': status,
            'response_excerpt': response_excerpt[:500],
            'route_binding_artifact': relpath(ROUTE),
            'artifact_count': len(uploaded_artifacts)
        },
        'confirmation_proof': {
            'confirmed_at_utc': datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace('+00:00', 'Z'),
            'archive_sha256': digest,
            'uploaded_bytes': size,
            'response_headers': headers
        },
        'uploaded_artifacts': uploaded_artifacts
    })
    write_json(RETIRE, {
        'retired_packet_id': packet_id,
        'retirement_reason': 'Route-backed backup confirmation completed for the recovered NASA budget workbook.',
        'post_confirmation_status': 'retired'
    })
    write_json(LEDGER, {
        'final_output': OUTPUT.name,
        'verification_checks': verification,
        'route_confirmation_artifact': relpath(DEST),
        'packet_retirement_artifact': relpath(RETIRE),
        'completion_status': 'complete'
    })


if __name__ == '__main__':
    main()
